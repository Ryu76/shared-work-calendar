import os
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

_tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
_tmp.close()
_upload_dir = tempfile.mkdtemp(prefix="shared-calendar-images-")
BOOTSTRAP_PASSWORD = "bootstrap-admin-password-123"
os.environ["WORK_CALENDAR_DB"] = _tmp.name
os.environ["WORK_CALENDAR_UPLOAD_DIR"] = _upload_dir
os.environ["WORK_CALENDAR_INITIAL_ADMIN_PASSWORD"] = BOOTSTRAP_PASSWORD
from fastapi.testclient import TestClient
import main

# Test discovery can import acceptance coverage before this module, leaving the
# cached application module pointed at that test's restored product DB path.
# Bind the already-imported module explicitly so this suite always uses its own
# isolated database and private-upload directory regardless of import order.
main.DB_PATH = Path(_tmp.name)
main.UPLOAD_DIR = Path(_upload_dir)
main.SESSIONS.clear()
client = TestClient(main.app)
ADMIN_PASSWORD = "admin-password-123"


class CalendarApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        main.init_db()
        initial = client.post("/api/auth/login", json={"login_id": "admin", "password": BOOTSTRAP_PASSWORD})
        assert initial.status_code == 200, initial.text
        changed = client.post(
            "/api/me/password",
            headers={"Authorization": "Bearer " + initial.json()["token"]},
            json={"current_password": BOOTSTRAP_PASSWORD, "new_password": ADMIN_PASSWORD},
        )
        assert changed.status_code == 200, changed.text

    def login(self, user, password):
        response = client.post("/api/auth/login", json={"login_id": user, "password": password})
        self.assertEqual(response.status_code, 200, response.text)
        return {"Authorization": "Bearer " + response.json()["token"]}, response.json()["user"]

    def admin(self):
        return self.login("admin", ADMIN_PASSWORD)[0]

    def create_member(self, admin, login_id):
        response = client.post(
            "/api/admin/users", headers=admin,
            json={"login_id": login_id, "name": login_id, "role": "MEMBER", "temporary_password": "password-123"},
        )
        self.assertEqual(response.status_code, 201, response.text)
        return response.json()["id"]

    def test_first_login_requires_password_change_and_blocks_work_apis(self):
        admin = self.admin()
        self.create_member(admin, "password_member")
        member, me = self.login("password_member", "password-123")
        self.assertTrue(me["must_change_password"])
        blocked = client.get("/api/tasks", headers=member)
        self.assertEqual(blocked.status_code, 403, blocked.text)
        self.assertIn("비밀번호를 변경", blocked.json()["detail"])
        wrong = client.post("/api/me/password", headers=member, json={"current_password": "wrong-password", "new_password": "new-password-123"})
        self.assertEqual(wrong.status_code, 400, wrong.text)
        changed = client.post("/api/me/password", headers=member, json={"current_password": "password-123", "new_password": "new-password-123"})
        self.assertEqual(changed.status_code, 200, changed.text)
        self.assertFalse(changed.json()["must_change_password"])
        self.assertEqual(client.get("/api/tasks", headers=member).status_code, 200)
        self.assertEqual(client.post("/api/auth/login", json={"login_id": "password_member", "password": "password-123"}).status_code, 401)
        new_login, new_me = self.login("password_member", "new-password-123")
        self.assertFalse(new_me["must_change_password"])
        audit = client.get("/api/admin/audit-logs", headers=admin)
        self.assertEqual(audit.status_code, 200, audit.text)
        self.assertTrue(any(log["action"] == "PASSWORD_CHANGED" and log["target_id"] == str(new_me["id"]) for log in audit.json()))
        self.assertEqual(client.post("/api/auth/logout", headers=new_login).status_code, 200)

    def test_argon2id_is_used_for_new_passwords_and_legacy_pbkdf2_upgrades_on_login(self):
        admin = self.admin()
        temporary_password = "password-123"
        modern_id = self.create_member(admin, "argon2_new_member")
        with main.db() as c:
            modern_hash = c.execute("SELECT password_hash FROM users WHERE id=?", (modern_id,)).fetchone()["password_hash"]
        self.assertTrue(modern_hash.startswith("$argon2id$"))
        self.assertNotIn(temporary_password, modern_hash)
        self.assertTrue(main.verify_password(temporary_password, modern_hash))
        self.assertFalse(main.verify_password("wrong-password", modern_hash))

        # Simulate an existing local database created before the Argon2id migration.
        legacy_password = "legacy-password-123"
        with main.db() as c:
            legacy_id = c.execute(
                "INSERT INTO users(login_id,password_hash,name,role,is_active,must_change_password,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?)",
                ("legacy_pbkdf2_member", main.legacy_hash_password(legacy_password), "legacy", "MEMBER", 1, 1, main.now(), main.now()),
            ).lastrowid
        rejected = client.post("/api/auth/login", json={"login_id": "legacy_pbkdf2_member", "password": "wrong-password"})
        self.assertEqual(rejected.status_code, 401, rejected.text)
        with main.db() as c:
            self.assertFalse(c.execute("SELECT password_hash FROM users WHERE id=?", (legacy_id,)).fetchone()["password_hash"].startswith("$argon2id$"))

        migrated = client.post("/api/auth/login", json={"login_id": "legacy_pbkdf2_member", "password": legacy_password})
        self.assertEqual(migrated.status_code, 200, migrated.text)
        with main.db() as c:
            migrated_hash = c.execute("SELECT password_hash FROM users WHERE id=?", (legacy_id,)).fetchone()["password_hash"]
        self.assertTrue(migrated_hash.startswith("$argon2id$"))
        self.assertNotIn(legacy_password, migrated_hash)
        self.assertTrue(main.verify_password(legacy_password, migrated_hash))
        audit = client.get("/api/admin/audit-logs", headers=admin)
        self.assertEqual(audit.status_code, 200, audit.text)
        self.assertTrue(any(log["action"] == "PASSWORD_HASH_UPGRADED" and log["target_id"] == str(legacy_id) for log in audit.json()))

    def test_admin_user_lifecycle_validation_authorization_and_last_admin_protection(self):
        admin = self.admin()
        created_id = self.create_member(admin, "user_management_member")
        member, _ = self.login("user_management_member", "password-123")
        changed = client.post(
            "/api/me/password", headers=member,
            json={"current_password": "password-123", "new_password": "member-password-123"},
        )
        self.assertEqual(changed.status_code, 200, changed.text)

        # Management endpoints remain server-admin-only and reject whitespace-only input.
        self.assertEqual(
            client.patch(f"/api/admin/users/{created_id}", headers=member, json={"is_active": False}).status_code,
            403,
        )
        self.assertEqual(
            client.post("/api/admin/users", headers=admin, json={
                "login_id": "blank_name_member", "name": "   ", "role": "MEMBER", "temporary_password": "password-123",
            }).status_code,
            422,
        )
        self.assertEqual(
            client.patch(f"/api/admin/users/{created_id}", headers=admin, json={"name": "  "}).status_code,
            422,
        )

        promoted = client.patch(
            f"/api/admin/users/{created_id}", headers=admin,
            json={"name": "수정된 사용자", "role": "ADMIN", "is_active": False},
        )
        self.assertEqual(promoted.status_code, 200, promoted.text)
        self.assertEqual(promoted.json(), {"id": created_id, "name": "수정된 사용자", "role": "ADMIN", "is_active": False})
        reactivated = client.patch(
            f"/api/admin/users/{created_id}", headers=admin,
            json={"is_active": True, "password": "reset-password-123"},
        )
        self.assertEqual(reactivated.status_code, 200, reactivated.text)
        relogin = client.post("/api/auth/login", json={"login_id": "user_management_member", "password": "reset-password-123"})
        self.assertEqual(relogin.status_code, 200, relogin.text)
        self.assertTrue(relogin.json()["user"]["must_change_password"])
        # The temporary second administrator can be demoted while the bootstrap
        # administrator remains active; that leaves exactly one active admin.
        demoted = client.patch(f"/api/admin/users/{created_id}", headers=admin, json={"role": "MEMBER"})
        self.assertEqual(demoted.status_code, 200, demoted.text)

        admin_id = client.get("/api/me", headers=admin).json()["id"]
        protected = client.patch(f"/api/admin/users/{admin_id}", headers=admin, json={"role": "MEMBER"})
        self.assertEqual(protected.status_code, 409, protected.text)
        self.assertIn("마지막 활성 관리자", protected.json()["detail"])
        audit = client.get("/api/admin/audit-logs", headers=admin)
        self.assertEqual(audit.status_code, 200, audit.text)
        user_updates = [row for row in audit.json() if row["action"] == "USER_UPDATED" and row["target_id"] == str(created_id)]
        self.assertTrue(user_updates)
        self.assertNotIn("reset-password-123", str(user_updates))

    def test_new_database_requires_a_bootstrap_secret(self):
        previous = os.environ.pop("WORK_CALENDAR_INITIAL_ADMIN_PASSWORD", None)
        try:
            with tempfile.TemporaryDirectory() as temp_dir, patch.object(main, "DB_PATH", Path(temp_dir) / "fresh.db"):
                with self.assertRaisesRegex(RuntimeError, "WORK_CALENDAR_INITIAL_ADMIN_PASSWORD"):
                    main.init_db()
        finally:
            if previous is not None:
                os.environ["WORK_CALENDAR_INITIAL_ADMIN_PASSWORD"] = previous

    def test_member_cannot_see_other_members_task(self):
        admin = self.admin()
        a_id = self.create_member(admin, "member_a")
        b_id = self.create_member(admin, "member_b")
        a, _ = self.login("member_a", "password-123")
        b, _ = self.login("member_b", "password-123")
        for headers in (a, b):
            changed = client.post("/api/me/password", headers=headers, json={"current_password": "password-123", "new_password": "member-password-123"})
            self.assertEqual(changed.status_code, 200, changed.text)
        response = client.post("/api/tasks", headers=b, json={"assignee_ids": [b_id], "contact": {"name": "요청자", "phone": "010-1234-5678", "organization": "기관"}, "description": "B만 보는 업무", "status": "RECEIVED"})
        self.assertEqual(response.status_code, 201, response.text)
        task_id = response.json()["id"]
        hidden_list = client.get("/api/tasks", headers=a)
        self.assertEqual(hidden_list.status_code, 200, hidden_list.text)
        self.assertEqual(hidden_list.json()["items"], [])
        self.assertEqual(client.get(f"/api/tasks/{task_id}", headers=a).status_code, 403)
        self.assertNotEqual(a_id, b_id)

    def test_completed_status_sets_and_clears_completed_date(self):
        admin = self.admin()
        admin_id = client.get("/api/me", headers=admin).json()["id"]
        create = client.post("/api/tasks", headers=admin, json={"assignee_ids": [admin_id], "contact": {"name": "완료검증", "phone": "010-7777-7777", "organization": "검증기관"}, "description": "상태 규칙 검증", "status": "COMPLETED"})
        self.assertEqual(create.status_code, 201, create.text)
        self.assertIsNotNone(create.json()["completed_date"])
        update = client.patch(f"/api/tasks/{create.json()['id']}", headers=admin, json={"assignee_ids": [admin_id], "contact": {"name": "완료검증", "phone": "010-7777-7777", "organization": "검증기관"}, "description": "상태 규칙 검증", "status": "IN_PROGRESS"})
        self.assertEqual(update.status_code, 200, update.text)
        self.assertIsNone(update.json()["completed_date"])

    def test_automatic_task_dates_use_the_asia_seoul_calendar_date(self):
        """Automatic task dates must use the product's Asia/Seoul calendar day."""
        admin = self.admin()
        admin_id = client.get("/api/me", headers=admin).json()["id"]
        seoul_today = main.date(2031, 1, 2)
        with patch.object(main, "today_in_seoul", return_value=seoul_today):
            created = client.post(
                "/api/tasks",
                headers=admin,
                json={
                    "assignee_ids": [admin_id],
                    "contact": {"name": "서울 날짜 요청", "phone": "010-7777-8888", "organization": "서울 날짜 기관"},
                    "description": "서울 기준 자동 날짜 검증",
                    "status": "RECEIVED",
                },
            )
            self.assertEqual(created.status_code, 201, created.text)
            self.assertEqual(created.json()["received_date"], "2031-01-02")
            completed = client.patch(
                f"/api/tasks/{created.json()['id']}", headers=admin, json={"status": "COMPLETED"}
            )
        self.assertEqual(completed.status_code, 200, completed.text)
        self.assertEqual(completed.json()["completed_date"], "2031-01-02")

    def test_task_list_server_filters_pagination_and_calendar_month_validation(self):
        admin = self.admin()
        member_id = self.create_member(admin, "filter_member")
        member, _ = self.login("filter_member", "password-123")
        changed = client.post("/api/me/password", headers=member, json={"current_password": "password-123", "new_password": "member-password-123"})
        self.assertEqual(changed.status_code, 200, changed.text)
        admin_id = client.get("/api/me", headers=admin).json()["id"]

        created_ids = []
        for description, task_status, received_date in [
            ("필터 접수 업무", "RECEIVED", "2026-09-01"),
            ("필터 진행 업무", "IN_PROGRESS", "2026-09-02"),
            ("필터 처리 업무", "COMPLETED", "2026-09-03"),
        ]:
            response = client.post("/api/tasks", headers=admin, json={
                "assignee_ids": [member_id],
                "contact": {"name": "필터 요청자", "phone": "010-4444-5555", "organization": "필터 기관"},
                "description": description,
                "status": task_status,
                "received_date": received_date,
            })
            self.assertEqual(response.status_code, 201, response.text)
            created_ids.append(response.json()["id"])

        first_page = client.get("/api/tasks?status=IN_PROGRESS&q=필터&received_from=2026-09-01&received_to=2026-09-02&assignee_id=" + str(member_id) + "&page=1&page_size=1", headers=admin)
        self.assertEqual(first_page.status_code, 200, first_page.text)
        result = first_page.json()
        self.assertEqual(result["pagination"], {"page": 1, "page_size": 1, "total": 1, "total_pages": 1})
        self.assertEqual([row["id"] for row in result["items"]], [created_ids[1]])
        self.assertEqual(result["summary"], {"RECEIVED": 0, "IN_PROGRESS": 1, "COMPLETED": 0})

        pages = client.get("/api/tasks?q=필터&assignee_id=" + str(member_id) + "&page=2&page_size=2", headers=admin)
        self.assertEqual(pages.status_code, 200, pages.text)
        self.assertEqual(pages.json()["pagination"], {"page": 2, "page_size": 2, "total": 3, "total_pages": 2})
        self.assertEqual(len(pages.json()["items"]), 1)
        self.assertEqual(pages.json()["summary"], {"RECEIVED": 1, "IN_PROGRESS": 1, "COMPLETED": 1})

        self.assertEqual(client.get("/api/tasks?assignee_id=" + str(admin_id), headers=member).status_code, 403)
        own_filter = client.get("/api/tasks?assignee_id=" + str(member_id), headers=member)
        self.assertEqual(own_filter.status_code, 200, own_filter.text)
        self.assertTrue(all(member_id in [assignee["id"] for assignee in task["assignees"]] for task in own_filter.json()["items"]))
        self.assertEqual(client.get("/api/tasks?status=INVALID", headers=admin).status_code, 422)
        self.assertEqual(client.get("/api/tasks?page_size=101", headers=admin).status_code, 422)
        self.assertEqual(client.get("/api/tasks?received_from=2026-09-03&received_to=2026-09-01", headers=admin).status_code, 422)
        self.assertEqual(client.get("/api/calendar?month=2026-7", headers=admin).status_code, 422)
        self.assertEqual(client.get("/api/calendar?month=2026-13", headers=admin).status_code, 422)
        valid_calendar = client.get("/api/calendar?month=2026-09", headers=admin)
        self.assertEqual(valid_calendar.status_code, 200, valid_calendar.text)
        # The received-date operational calendar keeps all live statuses visible;
        # completed entries render with their dedicated status color in the UI.
        self.assertEqual({row["id"] for row in valid_calendar.json()}, set(created_ids))
        self.assertEqual({row["status"] for row in valid_calendar.json()}, {"RECEIVED", "IN_PROGRESS", "COMPLETED"})

    def test_statistics_scope_status_user_aggregation_and_validation(self):
        admin = self.admin()
        analyst_id = self.create_member(admin, "statistics_member")
        analyst, _ = self.login("statistics_member", "password-123")
        changed = client.post(
            "/api/me/password", headers=analyst,
            json={"current_password": "password-123", "new_password": "member-password-123"},
        )
        self.assertEqual(changed.status_code, 200, changed.text)
        admin_id = client.get("/api/me", headers=admin).json()["id"]
        for description, assignees, task_status, received_date in (
            ("통계 접수 업무", [analyst_id], "RECEIVED", "2032-03-10"),
            ("통계 공동 처리 업무", [analyst_id, admin_id], "IN_PROGRESS", "2032-03-11"),
            ("통계 관리자 완료 업무", [admin_id], "COMPLETED", "2032-03-12"),
        ):
            created = client.post("/api/tasks", headers=admin, json={
                "assignee_ids": assignees,
                "contact": {"name": "통계 요청자", "phone": "010-1111-2222", "organization": "통계 기관"},
                "description": description,
                "status": task_status,
                "received_date": received_date,
            })
            self.assertEqual(created.status_code, 201, created.text)

        report = client.get(
            "/api/statistics?received_from=2032-03-10&received_to=2032-03-11", headers=admin
        )
        self.assertEqual(report.status_code, 200, report.text)
        self.assertEqual(report.json()["period"], {"received_from": "2032-03-10", "received_to": "2032-03-11"})
        self.assertEqual(report.json()["total"], 2)
        self.assertEqual(report.json()["by_status"], {"RECEIVED": 1, "IN_PROGRESS": 1, "COMPLETED": 0})
        analyst_row = next(row for row in report.json()["by_user"] if row["id"] == analyst_id)
        self.assertEqual(analyst_row["total"], 2)
        self.assertEqual(analyst_row["by_status"], {"RECEIVED": 1, "IN_PROGRESS": 1, "COMPLETED": 0})

        # A member sees only their authorized status aggregate and their own user row;
        # direct filters cannot enumerate a different user's workload.
        member_report = client.get("/api/statistics", headers=analyst)
        self.assertEqual(member_report.status_code, 200, member_report.text)
        self.assertEqual(member_report.json()["by_status"], {"RECEIVED": 1, "IN_PROGRESS": 1, "COMPLETED": 0})
        self.assertEqual([row["id"] for row in member_report.json()["by_user"]], [analyst_id])
        self.assertEqual(client.get(f"/api/statistics?assignee_id={admin_id}", headers=analyst).status_code, 403)
        self.assertEqual(
            client.get("/api/statistics?received_from=2032-03-12&received_to=2032-03-10", headers=admin).status_code,
            422,
        )
        with main.db() as c:
            indexes = {row["name"] for row in c.execute("PRAGMA index_list('tasks')")}
        self.assertIn("idx_tasks_live_received_status", indexes)

    def test_due_dates_local_notification_outbox_preferences_and_authorization(self):
        admin = self.admin()
        member_id = self.create_member(admin, "due_notification_member")
        outsider_id = self.create_member(admin, "due_notification_outsider")
        member, _ = self.login("due_notification_member", "password-123")
        outsider, _ = self.login("due_notification_outsider", "password-123")
        for headers in (member, outsider):
            changed = client.post(
                "/api/me/password", headers=headers,
                json={"current_password": "password-123", "new_password": "member-password-123"},
            )
            self.assertEqual(changed.status_code, 200, changed.text)
        admin_id = client.get("/api/me", headers=admin).json()["id"]
        invalid = client.post("/api/tasks", headers=admin, json={
            "assignee_ids": [member_id],
            "contact": {"name": "마감 요청", "phone": "010-3000-4000", "organization": "마감 기관"},
            "description": "마감일이 접수일보다 빠른 업무", "received_date": "2032-04-11", "due_date": "2032-04-10",
        })
        self.assertEqual(invalid.status_code, 422, invalid.text)
        created = client.post("/api/tasks", headers=admin, json={
            "assignee_ids": [member_id],
            "contact": {"name": "마감 요청", "phone": "010-3000-4000", "organization": "마감 기관"},
            "description": "멤버만 받는 오늘 마감 업무", "received_date": "2032-04-01", "due_date": "2032-04-10",
        })
        self.assertEqual(created.status_code, 201, created.text)
        task_id = created.json()["id"]
        self.assertEqual(created.json()["due_date"], "2032-04-10")
        listed = client.get("/api/tasks?due_from=2032-04-10&due_to=2032-04-10", headers=admin)
        self.assertEqual(listed.status_code, 200, listed.text)
        self.assertIn(task_id, [task["id"] for task in listed.json()["items"]])
        self.assertEqual(client.get("/api/tasks?due_from=2032-04-11&due_to=2032-04-10", headers=admin).status_code, 422)

        with patch.object(main, "today_in_seoul", return_value=main.date(2032, 4, 10)):
            preferences = client.get("/api/me/notification-preferences", headers=member)
            self.assertEqual(preferences.status_code, 200, preferences.text)
            self.assertEqual(preferences.json()["delivery"], "IN_APP_OUTBOX_ONLY")
            notifications = client.get("/api/notifications", headers=member)
        self.assertEqual(notifications.status_code, 200, notifications.text)
        self.assertEqual(notifications.json()["delivery"], "IN_APP_OUTBOX_ONLY")
        self.assertEqual(len(notifications.json()["items"]), 1)
        notification = notifications.json()["items"][0]
        self.assertEqual(notification["task_id"], task_id)
        self.assertEqual(notification["kind"], "DUE_TODAY")
        self.assertIn("오늘 마감", notification["message"])
        self.assertEqual(client.get("/api/notifications", headers=outsider).json()["items"], [])
        self.assertEqual(client.post(f"/api/notifications/{notification['id']}/read", headers=outsider).status_code, 404)
        self.assertEqual(client.post(f"/api/notifications/{notification['id']}/read", headers=member).status_code, 200)

        # A local reminder includes task text, so changing that text must invalidate
        # the old outbox item and lazily create one with the current description.
        renamed = client.patch(
            f"/api/tasks/{task_id}", headers=admin,
            json={"description": "수정된 오늘 마감 업무"},
        )
        self.assertEqual(renamed.status_code, 200, renamed.text)
        with patch.object(main, "today_in_seoul", return_value=main.date(2032, 4, 10)):
            refreshed_notifications = client.get("/api/notifications", headers=member)
        self.assertEqual(refreshed_notifications.status_code, 200, refreshed_notifications.text)
        self.assertEqual(len(refreshed_notifications.json()["items"]), 1)
        self.assertIn("수정된 오늘 마감 업무", refreshed_notifications.json()["items"][0]["message"])
        self.assertFalse(refreshed_notifications.json()["items"][0]["is_read"])

        disabled = client.patch(
            "/api/me/notification-preferences", headers=member,
            json={"due_soon_enabled": False, "overdue_enabled": False},
        )
        self.assertEqual(disabled.status_code, 200, disabled.text)
        self.assertFalse(disabled.json()["due_soon_enabled"])
        self.assertEqual(client.get("/api/notifications", headers=member).json()["items"], [])
        audit = client.get("/api/admin/audit-logs", headers=admin).json()
        self.assertTrue(any(row["action"] == "NOTIFICATION_PREFERENCES_UPDATED" and row["target_id"] == str(member_id) for row in audit))

        # Updating a due date to null is intentional and clears it without touching
        # pre-existing SQLite task data; task history labels the field change.
        cleared = client.patch(f"/api/tasks/{task_id}", headers=admin, json={"due_date": None})
        self.assertEqual(cleared.status_code, 200, cleared.text)
        self.assertIsNone(cleared.json()["due_date"])
        history = client.get(f"/api/tasks/{task_id}/history", headers=admin).json()
        self.assertTrue(any("마감일" in row["changed_fields"] for row in history if row["action"] == "TASK_UPDATED"))
        self.assertNotEqual(admin_id, outsider_id)

    def test_due_date_column_migrates_a_preexisting_sqlite_tasks_table(self):
        previous = main.DB_PATH
        try:
            with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
                legacy_db = Path(temp_dir) / "legacy.db"
                try:
                    # Production persistence is SQLCipher-backed SQLite. Create a
                    # legacy-shaped encrypted database so the migration exercises
                    # a deployable pre-existing database rather than an unsupported
                    # plaintext file opened with an encryption key.
                    with main.connect(legacy_db) as connection:
                        connection.execute(
                            "CREATE TABLE tasks (id INTEGER PRIMARY KEY, requester_contact_id INTEGER NOT NULL, status TEXT NOT NULL, "
                            "description TEXT NOT NULL, received_date TEXT NOT NULL, completed_date TEXT, created_by INTEGER NOT NULL, "
                            "created_at TEXT NOT NULL, updated_at TEXT NOT NULL, deleted_at TEXT, deleted_by INTEGER)"
                        )
                        # A real legacy row must survive the additive Phase-2 migration,
                        # not merely its table shape.  The pre-migration table deliberately
                        # has no due_date column and no foreign-key constraints, as can occur
                        # in a pre-existing SQLite deployment.
                        connection.execute(
                            "INSERT INTO tasks(id, requester_contact_id, status, description, received_date, completed_date, "
                            "created_by, created_at, updated_at, deleted_at, deleted_by) VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                            (77, 12, "RECEIVED", "기존 SQLite 업무", "2030-01-02", None, 5, "2030-01-02T00:00:00+00:00", "2030-01-02T00:00:00+00:00", None, None),
                        )
                    main.DB_PATH = legacy_db
                    main.init_db()
                    with main.db() as c:
                        columns = {row["name"] for row in c.execute("PRAGMA table_info('tasks')")}
                        indexes = {row["name"] for row in c.execute("PRAGMA index_list('tasks')")}
                        tables = {row["name"] for row in c.execute("SELECT name FROM sqlite_master WHERE type='table'")}
                        attachment_indexes = {row["name"] for row in c.execute("PRAGMA index_list('task_attachments')")}
                        legacy_row = c.execute(
                            "SELECT id, description, received_date, due_date FROM tasks WHERE id=77"
                        ).fetchone()
                    self.assertIn("due_date", columns)
                    self.assertIn("idx_tasks_live_due_date", indexes)
                    self.assertIn("task_attachments", tables)
                    self.assertIn("task_image_edits", tables)
                    self.assertIn("idx_task_attachments_task_created", attachment_indexes)
                    self.assertEqual(dict(legacy_row), {
                        "id": 77, "description": "기존 SQLite 업무", "received_date": "2030-01-02", "due_date": None,
                    })
                finally:
                    # Restore the product test database before Windows cleans this
                    # temporary directory, avoiding a lingering SQLite file handle.
                    main.DB_PATH = previous
        finally:
            main.DB_PATH = previous

    def test_task_detail_update_history_and_deletion_authorization(self):
        admin = self.admin()
        creator_id = self.create_member(admin, "detail_creator")
        partner_id = self.create_member(admin, "detail_partner")
        self.create_member(admin, "detail_outside")
        creator, _ = self.login("detail_creator", "password-123")
        partner, _ = self.login("detail_partner", "password-123")
        outsider, _ = self.login("detail_outside", "password-123")
        for headers in (creator, partner, outsider):
            changed = client.post("/api/me/password", headers=headers, json={"current_password": "password-123", "new_password": "member-password-123"})
            self.assertEqual(changed.status_code, 200, changed.text)
        selectable = client.get("/api/assignees", headers=creator)
        self.assertEqual(selectable.status_code, 200, selectable.text)
        self.assertTrue(any(row["id"] == partner_id for row in selectable.json()))
        invalid = client.post("/api/tasks", headers=creator, json={"assignee_ids": [creator_id], "contact": {"name": "요청", "phone": "not-a-phone", "organization": "기관"}, "description": "잘못된 연락처", "status": "RECEIVED"})
        self.assertEqual(invalid.status_code, 422, invalid.text)
        created = client.post("/api/tasks", headers=creator, json={"assignee_ids": [creator_id, partner_id], "contact": {"name": "상세 요청자", "phone": "010-3333-4444", "organization": "상세 기관"}, "description": "상세 화면 검증 업무", "status": "RECEIVED", "received_date": "2026-07-15"})
        self.assertEqual(created.status_code, 201, created.text)
        task_id = created.json()["id"]
        self.assertEqual(client.get(f"/api/tasks/{task_id}", headers=partner).status_code, 200)
        self.assertEqual(client.get(f"/api/tasks/{task_id}", headers=outsider).status_code, 403)
        self.assertEqual(client.get(f"/api/tasks/{task_id}/history", headers=outsider).status_code, 403)
        patched = client.patch(f"/api/tasks/{task_id}", headers=partner, json={"description": "처리자가 수정한 업무 설명"})
        self.assertEqual(patched.status_code, 200, patched.text)
        self.assertEqual(patched.json()["received_date"], "2026-07-15")
        self.assertEqual(patched.json()["description"], "처리자가 수정한 업무 설명")
        # A participating member can edit task data, but cannot use a direct PATCH
        # request to move the received date owned by the task creator.
        received_date_bypass = client.patch(
            f"/api/tasks/{task_id}", headers=partner, json={"received_date": "2026-07-16"}
        )
        self.assertEqual(received_date_bypass.status_code, 403, received_date_bypass.text)
        self.assertIn("접수일", received_date_bypass.json()["detail"])
        creator_received_update = client.patch(
            f"/api/tasks/{task_id}", headers=creator, json={"received_date": "2026-07-16"}
        )
        self.assertEqual(creator_received_update.status_code, 200, creator_received_update.text)
        self.assertEqual(creator_received_update.json()["received_date"], "2026-07-16")
        history = client.get(f"/api/tasks/{task_id}/history", headers=creator)
        self.assertEqual(history.status_code, 200, history.text)
        update = next(row for row in history.json() if row["action"] == "TASK_UPDATED" and row["actor_name"] == "detail_partner")
        self.assertEqual(update["actor_name"], "detail_partner")
        self.assertIn("업무 설명", update["changed_fields"])
        self.assertEqual(client.delete(f"/api/tasks/{task_id}", headers=partner).status_code, 403)
        self.assertEqual(client.delete(f"/api/tasks/{task_id}", headers=admin).status_code, 200)
        self.assertEqual(client.get(f"/api/tasks/{task_id}", headers=admin).status_code, 404)

    def test_private_image_upload_validation_authorization_and_deletion(self):
        admin = self.admin()
        owner_id = self.create_member(admin, "image_owner")
        self.create_member(admin, "image_outsider")
        owner, _ = self.login("image_owner", "password-123")
        outsider, _ = self.login("image_outsider", "password-123")
        for headers in (owner, outsider):
            changed = client.post("/api/me/password", headers=headers, json={"current_password": "password-123", "new_password": "member-password-123"})
            self.assertEqual(changed.status_code, 200, changed.text)
        created = client.post("/api/tasks", headers=owner, json={
            "assignee_ids": [owner_id], "contact": {"name": "이미지 요청", "phone": "010-5555-6666", "organization": "이미지 기관"},
            "description": "비공개 이미지 업무", "status": "RECEIVED",
        })
        self.assertEqual(created.status_code, 201, created.text)
        task_id = created.json()["id"]
        png = b"\x89PNG\r\n\x1a\n" + b"valid-test-png"
        forbidden = client.post(f"/api/tasks/{task_id}/images", headers=outsider, files=[("images", ("secret.png", png, "image/png"))])
        self.assertEqual(forbidden.status_code, 403, forbidden.text)
        spoofed = client.post(f"/api/tasks/{task_id}/images", headers=owner, files=[("images", ("spoof.jpg", png, "image/jpeg"))])
        self.assertEqual(spoofed.status_code, 422, spoofed.text)
        uploaded = client.post(f"/api/tasks/{task_id}/images", headers=owner, files=[("images", ("secret.png", png, "image/png"))])
        self.assertEqual(uploaded.status_code, 201, uploaded.text)
        image = uploaded.json()["images"][0]
        self.assertNotIn("storage_key", image)
        detail = client.get(f"/api/tasks/{task_id}", headers=owner)
        self.assertEqual(detail.status_code, 200, detail.text)
        self.assertEqual(detail.json()["images"][0]["id"], image["id"])
        image_history = client.get(f"/api/tasks/{task_id}/history", headers=owner)
        self.assertEqual(image_history.status_code, 200, image_history.text)
        self.assertTrue(any(row["action"] == "TASK_IMAGE_UPLOADED" for row in image_history.json()))
        self.assertEqual(client.get(f"/api/tasks/{task_id}/images/{image['id']}", headers=outsider).status_code, 403)
        served = client.get(f"/api/tasks/{task_id}/images/{image['id']}", headers=owner)
        self.assertEqual(served.status_code, 200, served.text)
        self.assertEqual(served.content, png)
        self.assertEqual(served.headers["content-type"], "image/png")
        self.assertEqual(served.headers["x-content-type-options"], "nosniff")
        too_many = client.post(f"/api/tasks/{task_id}/images", headers=owner, files=[("images", (f"{i}.png", png, "image/png")) for i in range(5)])
        self.assertEqual(too_many.status_code, 422, too_many.text)
        self.assertEqual(client.delete(f"/api/tasks/{task_id}/images/{image['id']}", headers=outsider).status_code, 403)
        deleted = client.delete(f"/api/tasks/{task_id}/images/{image['id']}", headers=owner)
        self.assertEqual(deleted.status_code, 200, deleted.text)
        self.assertEqual(client.get(f"/api/tasks/{task_id}/images/{image['id']}", headers=owner).status_code, 404)
        audit = client.get("/api/admin/audit-logs", headers=admin).json()
        self.assertTrue(any(row["action"] == "TASK_IMAGE_UPLOADED" and row["target_id"] == str(task_id) for row in audit))
        self.assertTrue(any(row["action"] == "TASK_IMAGE_DELETED" and row["target_id"] == str(task_id) for row in audit))

    def test_private_image_rotation_is_authorized_non_destructive_and_audited(self):
        admin = self.admin()
        owner_id = self.create_member(admin, "image_editor_owner")
        self.create_member(admin, "image_editor_outsider")
        owner, _ = self.login("image_editor_owner", "password-123")
        outsider, _ = self.login("image_editor_outsider", "password-123")
        for headers in (owner, outsider):
            changed = client.post(
                "/api/me/password", headers=headers,
                json={"current_password": "password-123", "new_password": "member-password-123"},
            )
            self.assertEqual(changed.status_code, 200, changed.text)
        created = client.post("/api/tasks", headers=owner, json={
            "assignee_ids": [owner_id],
            "contact": {"name": "이미지 편집 요청", "phone": "010-7070-8080", "organization": "이미지 편집 기관"},
            "description": "비공개 이미지 회전 업무", "status": "RECEIVED",
        })
        self.assertEqual(created.status_code, 201, created.text)
        task_id = created.json()["id"]
        original_file = BytesIO()
        main.Image.new("RGB", (2, 3), (20, 40, 60)).save(original_file, format="PNG")
        original_bytes = original_file.getvalue()
        uploaded = client.post(
            f"/api/tasks/{task_id}/images", headers=owner,
            files=[("images", ("rotate.png", original_bytes, "image/png"))],
        )
        self.assertEqual(uploaded.status_code, 201, uploaded.text)
        image_id = uploaded.json()["images"][0]["id"]
        with main.db() as c:
            original_key = c.execute("SELECT storage_key FROM task_images WHERE id=?", (image_id,)).fetchone()["storage_key"]
        self.assertEqual(
            client.post(f"/api/tasks/{task_id}/images/{image_id}/edit", headers=outsider, json={"direction": "CLOCKWISE"}).status_code,
            403,
        )
        self.assertEqual(
            client.post(f"/api/tasks/{task_id}/images/{image_id}/edit", headers=owner, json={"direction": "UPSIDE_DOWN"}).status_code,
            422,
        )
        rotated = client.post(
            f"/api/tasks/{task_id}/images/{image_id}/edit", headers=owner, json={"direction": "CLOCKWISE"}
        )
        self.assertEqual(rotated.status_code, 200, rotated.text)
        self.assertTrue(rotated.json()["non_destructive"])
        self.assertEqual(rotated.json()["direction"], "CLOCKWISE")
        self.assertNotIn("storage_key", rotated.text)
        served = client.get(f"/api/tasks/{task_id}/images/{image_id}", headers=owner)
        self.assertEqual(served.status_code, 200, served.text)
        with main.Image.open(BytesIO(served.content)) as rendered:
            self.assertEqual(rendered.size, (3, 2))
        with main.db() as c:
            image_row = c.execute("SELECT storage_key FROM task_images WHERE id=?", (image_id,)).fetchone()
            revision = c.execute("SELECT * FROM task_image_edits WHERE image_id=?", (image_id,)).fetchone()
        self.assertNotEqual(image_row["storage_key"], original_key)
        self.assertEqual(revision["source_storage_key"], original_key)
        self.assertTrue(main.private_image_path(original_key).is_file())
        self.assertTrue(main.private_image_path(image_row["storage_key"]).is_file())
        history = client.get(f"/api/tasks/{task_id}/history", headers=owner)
        self.assertTrue(any(row["action"] == "TASK_IMAGE_EDITED" for row in history.json()))
        audit = client.get("/api/admin/audit-logs?action=TASK_IMAGE_EDITED", headers=admin)
        self.assertEqual(audit.status_code, 200, audit.text)
        self.assertTrue(any(row["target_id"] == str(task_id) for row in audit.json()))
        deleted = client.delete(f"/api/tasks/{task_id}/images/{image_id}", headers=owner)
        self.assertEqual(deleted.status_code, 200, deleted.text)
        self.assertFalse(main.private_image_path(original_key).exists())
        self.assertFalse(main.private_image_path(image_row["storage_key"]).exists())

    def test_private_document_video_attachment_validation_authorization_and_deletion(self):
        admin = self.admin()
        owner_id = self.create_member(admin, "attachment_owner")
        self.create_member(admin, "attachment_outsider")
        owner, _ = self.login("attachment_owner", "password-123")
        outsider, _ = self.login("attachment_outsider", "password-123")
        for headers in (owner, outsider):
            changed = client.post("/api/me/password", headers=headers, json={"current_password": "password-123", "new_password": "member-password-123"})
            self.assertEqual(changed.status_code, 200, changed.text)
        created = client.post("/api/tasks", headers=owner, json={
            "assignee_ids": [owner_id], "contact": {"name": "첨부 요청", "phone": "010-1212-3434", "organization": "첨부 기관"},
            "description": "비공개 문서와 동영상 업무", "status": "RECEIVED",
        })
        self.assertEqual(created.status_code, 201, created.text)
        task_id = created.json()["id"]
        pdf = b"%PDF-1.4\nphase-two-private-document\n%%EOF"
        mp4 = b"\x00\x00\x00\x18ftypisom\x00\x00\x00\x00isomiso2"
        self.assertEqual(
            client.post(f"/api/tasks/{task_id}/attachments", headers=outsider, files=[("attachments", ("secret.pdf", pdf, "application/pdf"))]).status_code,
            403,
        )
        self.assertEqual(
            client.post(f"/api/tasks/{task_id}/attachments", headers=owner, files=[("attachments", ("spoof.pdf", b"not-a-pdf", "application/pdf"))]).status_code,
            422,
        )
        self.assertEqual(
            client.post(f"/api/tasks/{task_id}/attachments", headers=owner, files=[("attachments", ("unsafe.exe", b"MZ", "application/octet-stream"))]).status_code,
            422,
        )
        uploaded = client.post(
            f"/api/tasks/{task_id}/attachments", headers=owner,
            files=[
                ("attachments", ("private-report.pdf", pdf, "application/pdf")),
                ("attachments", ("private-video.mp4", mp4, "video/mp4")),
            ],
        )
        self.assertEqual(uploaded.status_code, 201, uploaded.text)
        attachments = uploaded.json()["attachments"]
        self.assertEqual({item["attachment_type"] for item in attachments}, {"DOCUMENT", "VIDEO"})
        self.assertTrue(all("storage_key" not in item and "sha256" not in item for item in attachments))
        document = next(item for item in attachments if item["attachment_type"] == "DOCUMENT")
        detail = client.get(f"/api/tasks/{task_id}", headers=owner)
        self.assertEqual(detail.status_code, 200, detail.text)
        self.assertEqual({item["id"] for item in detail.json()["attachments"]}, {item["id"] for item in attachments})
        forbidden_download = client.get(f"/api/tasks/{task_id}/attachments/{document['id']}", headers=outsider)
        self.assertEqual(forbidden_download.status_code, 403, forbidden_download.text)
        served = client.get(f"/api/tasks/{task_id}/attachments/{document['id']}", headers=owner)
        self.assertEqual(served.status_code, 200, served.text)
        self.assertEqual(served.content, pdf)
        self.assertEqual(served.headers["cache-control"], "private, no-store")
        self.assertEqual(served.headers["x-content-type-options"], "nosniff")
        self.assertTrue(served.headers["content-disposition"].startswith("attachment;"))
        self.assertEqual(client.delete(f"/api/tasks/{task_id}/attachments/{document['id']}", headers=outsider).status_code, 403)
        deleted = client.delete(f"/api/tasks/{task_id}/attachments/{document['id']}", headers=owner)
        self.assertEqual(deleted.status_code, 200, deleted.text)
        self.assertEqual(client.get(f"/api/tasks/{task_id}/attachments/{document['id']}", headers=owner).status_code, 404)
        history = client.get(f"/api/tasks/{task_id}/history", headers=owner).json()
        self.assertTrue(any(row["action"] == "TASK_ATTACHMENT_UPLOADED" for row in history))
        self.assertTrue(any(row["action"] == "TASK_ATTACHMENT_DELETED" for row in history))

    def test_completed_date_report_calendar_uses_task_scope_and_month_validation(self):
        admin = self.admin()
        member_id = self.create_member(admin, "completed_report_member")
        outsider_id = self.create_member(admin, "completed_report_outsider")
        member, _ = self.login("completed_report_member", "password-123")
        changed = client.post(
            "/api/me/password", headers=member,
            json={"current_password": "password-123", "new_password": "member-password-123"},
        )
        self.assertEqual(changed.status_code, 200, changed.text)
        visible = client.post("/api/tasks", headers=admin, json={
            "assignee_ids": [member_id],
            "contact": {"name": "처리일 리포트 요청", "phone": "010-2121-3434", "organization": "처리일 리포트 기관"},
            "description": "멤버 처리일 리포트 업무", "status": "COMPLETED",
            "received_date": "2034-04-01", "completed_date": "2034-04-12",
        })
        self.assertEqual(visible.status_code, 201, visible.text)
        hidden = client.post("/api/tasks", headers=admin, json={
            "assignee_ids": [outsider_id],
            "contact": {"name": "비공개 처리일 요청", "phone": "010-4545-6767", "organization": "비공개 처리일 기관"},
            "description": "다른 사용자의 처리일 리포트 업무", "status": "COMPLETED",
            "received_date": "2034-04-01", "completed_date": "2034-04-12",
        })
        self.assertEqual(hidden.status_code, 201, hidden.text)
        other_month = client.post("/api/tasks", headers=admin, json={
            "assignee_ids": [member_id],
            "contact": {"name": "다른 달 처리일 요청", "phone": "010-5656-7878", "organization": "다른 달 기관"},
            "description": "다른 달 완료 업무", "status": "COMPLETED",
            "received_date": "2034-04-01", "completed_date": "2034-05-01",
        })
        self.assertEqual(other_month.status_code, 201, other_month.text)
        report = client.get("/api/completed-calendar?month=2034-04", headers=member)
        self.assertEqual(report.status_code, 200, report.text)
        self.assertEqual([task["id"] for task in report.json()], [visible.json()["id"]])
        self.assertTrue(all(task["status"] == "COMPLETED" and task["completed_date"] == "2034-04-12" for task in report.json()))
        admin_report = client.get("/api/completed-calendar?month=2034-04", headers=admin)
        self.assertEqual(admin_report.status_code, 200, admin_report.text)
        self.assertEqual({task["id"] for task in admin_report.json()}, {visible.json()["id"], hidden.json()["id"]})
        self.assertEqual(client.get("/api/completed-calendar?month=2034-4", headers=admin).status_code, 422)
        self.assertEqual(client.get("/api/completed-calendar?month=2034-13", headers=admin).status_code, 422)

    def test_excel_export_matches_authorized_filtered_scope_and_is_spreadsheet_safe(self):
        admin = self.admin()
        member_id = self.create_member(admin, "excel_export_member")
        outsider_id = self.create_member(admin, "excel_export_outsider")
        member, _ = self.login("excel_export_member", "password-123")
        outsider, _ = self.login("excel_export_outsider", "password-123")
        for headers in (member, outsider):
            changed = client.post(
                "/api/me/password", headers=headers,
                json={"current_password": "password-123", "new_password": "member-password-123"},
            )
            self.assertEqual(changed.status_code, 200, changed.text)
        allowed = client.post("/api/tasks", headers=admin, json={
            "assignee_ids": [member_id],
            "contact": {"name": "엑셀 요청", "phone": "010-7878-8989", "organization": "엑셀 기관"},
            "description": "=수식으로 해석되면 안 되는 업무", "status": "IN_PROGRESS",
            "received_date": "2033-01-10", "due_date": "2033-01-15",
        })
        self.assertEqual(allowed.status_code, 201, allowed.text)
        hidden = client.post("/api/tasks", headers=admin, json={
            "assignee_ids": [outsider_id],
            "contact": {"name": "비공개 엑셀 요청", "phone": "010-9090-1010", "organization": "비공개 기관"},
            "description": "내보내기에 포함되면 안 되는 업무", "status": "RECEIVED",
            "received_date": "2033-01-10",
        })
        self.assertEqual(hidden.status_code, 201, hidden.text)

        exported = client.get(
            "/api/tasks/export.xlsx?status=IN_PROGRESS&received_from=2033-01-10&received_to=2033-01-10",
            headers=member,
        )
        self.assertEqual(exported.status_code, 200, exported.text)
        self.assertEqual(exported.headers["content-type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertEqual(exported.headers["cache-control"], "private, no-store")
        self.assertEqual(exported.headers["x-content-type-options"], "nosniff")
        self.assertIn("attachment;", exported.headers["content-disposition"])
        book = load_workbook(BytesIO(exported.content), data_only=False)
        self.assertEqual(book.sheetnames, ["업무 목록", "내보내기 정보"])
        sheet = book["업무 목록"]
        self.assertEqual([cell.value for cell in sheet[1]], [
            "업무 ID", "상태", "업무 설명", "요청 담당자", "연락처", "소속 기관", "처리자",
            "접수일", "마감일", "처리일", "생성자", "생성 시각", "수정 시각",
        ])
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(2, 1).value, allowed.json()["id"])
        self.assertEqual(sheet.cell(2, 2).value, "진행중")
        self.assertEqual(sheet.cell(2, 3).value, "'=수식으로 해석되면 안 되는 업무")
        self.assertEqual(sheet.cell(2, 3).data_type, "s")
        self.assertNotIn("비공개", "\n".join(str(cell.value) for row in sheet.iter_rows() for cell in row))
        self.assertEqual(book["내보내기 정보"]["B4"].value, 1)
        self.assertEqual(
            client.get(f"/api/tasks/export.xlsx?assignee_id={outsider_id}", headers=member).status_code,
            403,
        )
        self.assertEqual(
            client.get("/api/tasks/export.xlsx?due_from=2033-01-16&due_to=2033-01-15", headers=admin).status_code,
            422,
        )
        audit = client.get("/api/admin/audit-logs?action=TASKS_EXPORTED", headers=admin)
        self.assertEqual(audit.status_code, 200, audit.text)
        self.assertTrue(any(row["actor_user_id"] == member_id and row["after"]["count"] == 1 for row in audit.json()))

    def test_request_master_management_authorization_lifecycle_and_audit(self):
        admin = self.admin()
        self.create_member(admin, "master_member")
        member, _ = self.login("master_member", "password-123")
        changed = client.post("/api/me/password", headers=member, json={"current_password": "password-123", "new_password": "member-password-123"})
        self.assertEqual(changed.status_code, 200, changed.text)
        direct_organization = client.post("/api/organizations", headers=member, json={"name": "직접 등록 기관"})
        self.assertEqual(direct_organization.status_code, 201, direct_organization.text)
        self.assertTrue(direct_organization.json()["is_active"])
        duplicate_organization = client.post("/api/organizations", headers=member, json={"name": "직접 등록 기관"})
        self.assertEqual(duplicate_organization.status_code, 201, duplicate_organization.text)
        self.assertEqual(duplicate_organization.json()["id"], direct_organization.json()["id"])
        created = client.post("/api/request-contacts", headers=member, json={"name": "마스터 담당자", "phone": "010-8888-9999", "organization": "마스터 기관"})
        self.assertEqual(created.status_code, 201, created.text)
        contact = created.json()
        self.assertTrue(contact["is_active"])
        self.assertEqual(contact["organization"], "마스터 기관")
        organizations = client.get("/api/organizations?q=마스터", headers=member)
        self.assertEqual(organizations.status_code, 200, organizations.text)
        self.assertTrue(any(row["name"] == "마스터 기관" and row["is_active"] for row in organizations.json()))
        visible = client.get("/api/request-contacts?q=마스터", headers=member)
        self.assertEqual(visible.status_code, 200, visible.text)
        self.assertTrue(any(row["id"] == contact["id"] for row in visible.json()))
        self.assertEqual(client.get("/api/admin/organizations", headers=member).status_code, 403)
        admin_contacts = client.get("/api/admin/request-contacts?q=마스터", headers=admin)
        self.assertEqual(admin_contacts.status_code, 200, admin_contacts.text)
        self.assertTrue(any(row["id"] == contact["id"] for row in admin_contacts.json()))
        self.assertEqual(client.patch(f"/api/admin/request-contacts/{contact['id']}", headers=member, json={"is_active": False}).status_code, 403)
        org = next(row for row in client.get("/api/admin/organizations?q=마스터", headers=admin).json() if row["name"] == "마스터 기관")
        renamed = client.patch(f"/api/admin/organizations/{org['id']}", headers=admin, json={"name": "수정 마스터 기관"})
        self.assertEqual(renamed.status_code, 200, renamed.text)
        updated = client.patch(f"/api/admin/request-contacts/{contact['id']}", headers=admin, json={"name": "수정 담당자", "phone": "010-8888-9999", "organization_id": org["id"], "is_active": False})
        self.assertEqual(updated.status_code, 200, updated.text)
        self.assertFalse(updated.json()["is_active"])
        self.assertEqual(client.get("/api/request-contacts?q=수정", headers=member).json(), [])
        audit = client.get("/api/admin/audit-logs", headers=admin).json()
        self.assertTrue(any(row["action"] == "ORGANIZATION_UPDATED" and row["target_id"] == str(org["id"]) for row in audit))
        self.assertTrue(any(row["action"] == "REQUEST_CONTACT_UPDATED" and row["target_id"] == str(contact["id"]) for row in audit))

    def test_audit_log_filters_authorization_validation_and_sensitive_redaction(self):
        admin = self.admin()
        admin_id = client.get("/api/me", headers=admin).json()["id"]
        member_id = self.create_member(admin, "audit_filter_member")
        member, _ = self.login("audit_filter_member", "password-123")
        changed = client.post(
            "/api/me/password", headers=member,
            json={"current_password": "password-123", "new_password": "member-password-123"},
        )
        self.assertEqual(changed.status_code, 200, changed.text)
        created_day = main.now()[:10]
        with main.db() as c:
            main.audit(
                c, admin_id, "FILTER_TEST_ALPHA", "TEST", "filter-alpha",
                before={"visible": "before", "password": "sensitive-test-value"},
                after={"visible": "after", "storage_key": "private/test.png"},
            )
            main.audit(c, member_id, "FILTER_TEST_BETA", "TEST", "filter-beta", after={"visible": "other"})
            c.execute(
                "INSERT INTO audit_logs(actor_user_id,action,target_type,target_id,created_at) VALUES(?,?,?,?,?)",
                (admin_id, "FILTER_TEST_OLD", "TEST", "filter-old", "2000-01-01T00:00:00+00:00"),
            )

        # Audit records are administrator-only, including direct API calls by members.
        self.assertEqual(client.get("/api/admin/audit-logs", headers=member).status_code, 403)
        filtered = client.get(
            "/api/admin/audit-logs",
            headers=admin,
            params={
                "action": "FILTER_TEST_ALPHA", "actor_user_id": admin_id,
                "created_from": created_day, "created_to": created_day,
            },
        )
        self.assertEqual(filtered.status_code, 200, filtered.text)
        self.assertEqual(len(filtered.json()), 1)
        entry = filtered.json()[0]
        self.assertEqual(entry["action"], "FILTER_TEST_ALPHA")
        self.assertEqual(entry["actor_user_id"], admin_id)
        self.assertEqual(entry["before"]["visible"], "before")
        self.assertEqual(entry["before"]["password"], "[REDACTED]")
        self.assertEqual(entry["after"]["storage_key"], "[REDACTED]")
        self.assertNotIn("before_data", entry)
        self.assertNotIn("after_data", entry)

        by_actor = client.get("/api/admin/audit-logs", headers=admin, params={"actor_user_id": member_id})
        self.assertEqual(by_actor.status_code, 200, by_actor.text)
        self.assertTrue(by_actor.json())
        self.assertTrue(all(row["actor_user_id"] == member_id for row in by_actor.json()))
        outside_range = client.get(
            "/api/admin/audit-logs", headers=admin,
            params={"action": "FILTER_TEST_ALPHA", "created_from": "2099-01-01", "created_to": "2099-01-02"},
        )
        self.assertEqual(outside_range.status_code, 200, outside_range.text)
        self.assertEqual(outside_range.json(), [])
        self.assertEqual(client.get("/api/admin/audit-logs?action=invalid", headers=admin).status_code, 422)
        self.assertEqual(client.get("/api/admin/audit-logs?actor_user_id=0", headers=admin).status_code, 422)
        self.assertEqual(
            client.get("/api/admin/audit-logs?created_from=2026-02-02&created_to=2026-02-01", headers=admin).status_code,
            422,
        )


if __name__ == "__main__":
    unittest.main()
