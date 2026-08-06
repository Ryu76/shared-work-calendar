from __future__ import annotations

import hashlib
import hmac
import os
import re
import secrets
import uuid
import zipfile
import warnings
from io import BytesIO
from contextlib import contextmanager
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from zoneinfo import ZoneInfo

from argon2 import PasswordHasher
from argon2.exceptions import InvalidHashError, VerificationError
from fastapi import Depends, File, FastAPI, HTTPException, Query, Request, UploadFile, status
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, UnidentifiedImageError


# Persistence is deliberately isolated: database.py enforces SQLCipher encryption
# and gets its 256-bit key from the user's Windows DPAPI-protected key file.
try:  # Imported as ``app.main`` with backend/ on sys.path.
    from core.config import load_settings
    from core.logging import configure_logging
    from database import DB_PATH, connect, sqlite3
    from schemas.requests import (
        ContactIn, ContactPatch, NotificationPreferencesIn,
        OrganizationIn, OrganizationPatch, TaskIn, TaskUpdateIn,
    )
    from routers.auth import create_router as create_auth_router
    from routers.tasks import create_router as create_task_router
    from routers.users import create_router as create_users_router
    from services.auth_service import AuthService
    from services.task_service import TaskService
    from services.user_service import UserService
except ModuleNotFoundError:  # Imported as ``backend.app.main`` from repository root.
    from backend.core.config import load_settings
    from backend.core.logging import configure_logging
    from backend.database import DB_PATH, connect, sqlite3
    from backend.schemas.requests import (
        ContactIn, ContactPatch, NotificationPreferencesIn,
        OrganizationIn, OrganizationPatch, TaskIn, TaskUpdateIn,
    )
    from backend.routers.auth import create_router as create_auth_router
    from backend.routers.tasks import create_router as create_task_router
    from backend.routers.users import create_router as create_users_router
    from backend.services.auth_service import AuthService
    from backend.services.task_service import TaskService
    from backend.services.user_service import UserService

# Keep the original backend-root configuration boundary after moving this module
# beneath ``backend/app``. This preserves static and private-upload locations.
BASE_DIR = Path(__file__).resolve().parent.parent
SETTINGS = load_settings(BASE_DIR)
LOGGER = configure_logging(SETTINGS.log_level)
STATIC_DIR = SETTINGS.static_dir
# This directory is deliberately not mounted as static content. Files are served only
# after task-level authorization succeeds.
UPLOAD_DIR = SETTINGS.upload_dir
MAX_IMAGE_COUNT = 5
MAX_IMAGE_BYTES = 10 * 1024 * 1024
# Editing is intentionally bounded independently of compressed file size so a
# decompression-bomb image cannot exhaust the local application host.
MAX_IMAGE_PIXELS = 40_000_000
IMAGE_SIGNATURES = {
    "image/jpeg": lambda data: data.startswith(b"\xff\xd8\xff"),
    "image/png": lambda data: data.startswith(b"\x89PNG\r\n\x1a\n"),
    "image/webp": lambda data: len(data) >= 12 and data[:4] == b"RIFF" and data[8:12] == b"WEBP",
}
IMAGE_SUFFIXES = {"image/jpeg": ".jpg", "image/png": ".png", "image/webp": ".webp"}
MAX_ATTACHMENT_COUNT = 10
MAX_DOCUMENT_BYTES = 25 * 1024 * 1024
MAX_VIDEO_BYTES = 100 * 1024 * 1024
# Non-image files are limited to formats whose declared MIME type can be checked
# against their bytes. Filename extensions and browser MIME claims are never trusted.
ATTACHMENT_RULES = {
    "application/pdf": ("DOCUMENT", ".pdf", MAX_DOCUMENT_BYTES),
    "text/plain": ("DOCUMENT", ".txt", MAX_DOCUMENT_BYTES),
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ("DOCUMENT", ".docx", MAX_DOCUMENT_BYTES),
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ("DOCUMENT", ".xlsx", MAX_DOCUMENT_BYTES),
    "video/mp4": ("VIDEO", ".mp4", MAX_VIDEO_BYTES),
    "video/webm": ("VIDEO", ".webm", MAX_VIDEO_BYTES),
}
app = FastAPI(title="Shared Work Calendar MVP", version="1.0.0")
app.mount("/assets", StaticFiles(directory=STATIC_DIR), name="assets")
SESSIONS: dict[str, tuple[int, datetime]] = {}
# Argon2id is deliberately the default (argon2-cffi's PasswordHasher type).  These
# parameters meet the interactive-password baseline while keeping the values in the
# encoded hash, so stronger future settings can trigger a transparent rehash.
PASSWORD_HASHER = PasswordHasher(time_cost=3, memory_cost=65536, parallelism=2, hash_len=32, salt_len=16)
SEOUL_TIMEZONE = ZoneInfo("Asia/Seoul")


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def today_in_seoul() -> date:
    """Return the calendar date used by the product's received/completed fields."""
    return datetime.now(SEOUL_TIMEZONE).date()


def legacy_hash_password(password: str, salt: bytes | None = None) -> str:
    """PBKDF2 format used only to verify and migrate pre-Argon2 local data."""
    salt = salt or secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), salt, 210_000)
    return f"{salt.hex()}${digest.hex()}"


def hash_password(password: str) -> str:
    """Create a non-reversible Argon2id password hash for every new password."""
    return PASSWORD_HASHER.hash(password)


def verify_password(password: str, encoded: str) -> bool:
    """Verify Argon2id hashes and the previous PBKDF2 format during migration."""
    try:
        if encoded.startswith("$argon2"):
            return PASSWORD_HASHER.verify(encoded, password)
        salt_hex, digest_hex = encoded.split("$", 1)
        actual = legacy_hash_password(password, bytes.fromhex(salt_hex)).split("$", 1)[1]
        return hmac.compare_digest(actual, digest_hex)
    except (InvalidHashError, VerificationError, ValueError, AttributeError, TypeError):
        return False


def password_hash_needs_upgrade(encoded: str) -> bool:
    """Identify PBKDF2 rows and Argon2 rows made with superseded parameters."""
    if not encoded.startswith("$argon2"):
        return True
    try:
        return PASSWORD_HASHER.check_needs_rehash(encoded)
    except (InvalidHashError, VerificationError):
        # Login verification will reject malformed hashes before this function runs.
        return False


@contextmanager
def db():
    """Compatibility boundary for routes and isolated tests; encryption lives in database.py."""
    connection = connect(DB_PATH)
    try:
        yield connection
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


def init_db() -> None:
    with db() as c:
        c.executescript(
            """
            CREATE TABLE IF NOT EXISTS users (
              id INTEGER PRIMARY KEY, login_id TEXT UNIQUE NOT NULL,
              password_hash TEXT NOT NULL, name TEXT NOT NULL,
              role TEXT NOT NULL CHECK(role IN ('ADMIN','MEMBER')),
              is_active INTEGER NOT NULL DEFAULT 1,
              must_change_password INTEGER NOT NULL DEFAULT 1,
              created_at TEXT NOT NULL, updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS organizations (
              id INTEGER PRIMARY KEY, name TEXT UNIQUE NOT NULL,
              is_active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL, updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS request_contacts (
              id INTEGER PRIMARY KEY, name TEXT NOT NULL, phone TEXT NOT NULL,
              organization_id INTEGER NOT NULL REFERENCES organizations(id),
              is_active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL, updated_at TEXT NOT NULL,
              UNIQUE(name, phone, organization_id)
            );
            CREATE TABLE IF NOT EXISTS tasks (
              id INTEGER PRIMARY KEY, requester_contact_id INTEGER NOT NULL REFERENCES request_contacts(id),
              status TEXT NOT NULL CHECK(status IN ('RECEIVED','IN_PROGRESS','COMPLETED')),
              description TEXT NOT NULL, received_date TEXT NOT NULL, completed_date TEXT, due_date TEXT,
              created_by INTEGER NOT NULL REFERENCES users(id), created_at TEXT NOT NULL, updated_at TEXT NOT NULL,
              deleted_at TEXT, deleted_by INTEGER REFERENCES users(id),
              CHECK((status = 'COMPLETED' AND completed_date IS NOT NULL) OR (status != 'COMPLETED' AND completed_date IS NULL))
            );
            CREATE TABLE IF NOT EXISTS task_assignees (
              task_id INTEGER NOT NULL REFERENCES tasks(id) ON DELETE CASCADE,
              user_id INTEGER NOT NULL REFERENCES users(id), assigned_at TEXT NOT NULL,
              PRIMARY KEY(task_id,user_id)
            );
            CREATE TABLE IF NOT EXISTS task_images (
              id INTEGER PRIMARY KEY, task_id INTEGER NOT NULL REFERENCES tasks(id) ON DELETE CASCADE,
              storage_key TEXT UNIQUE NOT NULL, original_filename TEXT NOT NULL, mime_type TEXT NOT NULL,
              size_bytes INTEGER NOT NULL CHECK(size_bytes > 0 AND size_bytes <= 10485760),
              sort_order INTEGER NOT NULL, uploaded_by INTEGER NOT NULL REFERENCES users(id), created_at TEXT NOT NULL
            );
            CREATE TRIGGER IF NOT EXISTS task_images_max_count
            BEFORE INSERT ON task_images
            WHEN (SELECT COUNT(*) FROM task_images WHERE task_id=NEW.task_id) >= 5
            BEGIN SELECT RAISE(ABORT, '업무당 이미지는 최대 5개입니다.'); END;
            -- Image edits are non-destructive revisions. The current image row is
            -- updated to a private derivative while prior source files stay tracked
            -- until the image itself is intentionally deleted.
            CREATE TABLE IF NOT EXISTS task_image_edits (
              id INTEGER PRIMARY KEY, task_id INTEGER NOT NULL REFERENCES tasks(id) ON DELETE CASCADE,
              image_id INTEGER NOT NULL REFERENCES task_images(id) ON DELETE CASCADE,
              source_storage_key TEXT NOT NULL, result_storage_key TEXT UNIQUE NOT NULL,
              direction TEXT NOT NULL CHECK(direction IN ('CLOCKWISE','COUNTERCLOCKWISE')),
              edited_by INTEGER NOT NULL REFERENCES users(id), created_at TEXT NOT NULL
            );
            CREATE INDEX IF NOT EXISTS idx_task_image_edits_image_created
              ON task_image_edits(image_id, created_at, id);
            -- Kept separate from legacy task_images so existing SQLite image rows
            -- and their private storage paths survive the Phase 2 expansion intact.
            CREATE TABLE IF NOT EXISTS task_attachments (
              id INTEGER PRIMARY KEY, task_id INTEGER NOT NULL REFERENCES tasks(id) ON DELETE CASCADE,
              storage_key TEXT UNIQUE NOT NULL, original_filename TEXT NOT NULL, mime_type TEXT NOT NULL,
              attachment_type TEXT NOT NULL CHECK(attachment_type IN ('DOCUMENT','VIDEO')),
              size_bytes INTEGER NOT NULL CHECK(size_bytes > 0 AND size_bytes <= 104857600),
              sha256 TEXT NOT NULL, uploaded_by INTEGER NOT NULL REFERENCES users(id), created_at TEXT NOT NULL
            );
            CREATE TRIGGER IF NOT EXISTS task_attachments_max_count
            BEFORE INSERT ON task_attachments
            WHEN (SELECT COUNT(*) FROM task_attachments WHERE task_id=NEW.task_id) >= 10
            BEGIN SELECT RAISE(ABORT, '업무당 문서/동영상 첨부는 최대 10개입니다.'); END;
            CREATE TABLE IF NOT EXISTS audit_logs (
              id INTEGER PRIMARY KEY, actor_user_id INTEGER REFERENCES users(id), action TEXT NOT NULL,
              target_type TEXT NOT NULL, target_id TEXT NOT NULL, before_data TEXT, after_data TEXT,
              ip_address TEXT, created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS notification_preferences (
              user_id INTEGER PRIMARY KEY REFERENCES users(id) ON DELETE CASCADE,
              due_soon_enabled INTEGER NOT NULL DEFAULT 1 CHECK(due_soon_enabled IN (0,1)),
              overdue_enabled INTEGER NOT NULL DEFAULT 1 CHECK(overdue_enabled IN (0,1)),
              updated_at TEXT NOT NULL
            );
            -- This is a local in-app outbox only.  No external provider address,
            -- credential, or delivery worker is stored or invoked by the application.
            CREATE TABLE IF NOT EXISTS notification_outbox (
              id INTEGER PRIMARY KEY,
              user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
              task_id INTEGER NOT NULL REFERENCES tasks(id) ON DELETE CASCADE,
              kind TEXT NOT NULL CHECK(kind IN ('DUE_SOON','DUE_TODAY','OVERDUE')),
              due_date TEXT NOT NULL,
              message TEXT NOT NULL,
              read_at TEXT,
              created_at TEXT NOT NULL,
              UNIQUE(user_id, task_id, kind, due_date)
            );
            -- Non-destructive Phase 2 migration: safe for populated SQLite DBs.
            CREATE INDEX IF NOT EXISTS idx_tasks_live_received_status
              ON tasks(deleted_at, received_date, status);
            CREATE INDEX IF NOT EXISTS idx_task_assignees_user_task
              ON task_assignees(user_id, task_id);
            CREATE INDEX IF NOT EXISTS idx_notification_outbox_user_read
              ON notification_outbox(user_id, read_at, id DESC);
            CREATE INDEX IF NOT EXISTS idx_task_attachments_task_created
              ON task_attachments(task_id, created_at, id);
            """
        )
        # `CREATE TABLE IF NOT EXISTS` does not add columns to a pre-Phase-2
        # database.  This idempotent ALTER preserves every existing task row.
        task_columns = {row["name"] for row in c.execute("PRAGMA table_info('tasks')")}
        if "due_date" not in task_columns:
            c.execute("ALTER TABLE tasks ADD COLUMN due_date TEXT")
        c.execute("CREATE INDEX IF NOT EXISTS idx_tasks_live_due_date ON tasks(deleted_at, due_date, status)")
        if not c.execute("SELECT 1 FROM users WHERE login_id='admin'").fetchone():
            bootstrap_password = os.getenv("WORK_CALENDAR_INITIAL_ADMIN_PASSWORD")
            if not bootstrap_password or len(bootstrap_password) < 8:
                raise RuntimeError(
                    "새 데이터베이스를 초기화하려면 8자 이상의 WORK_CALENDAR_INITIAL_ADMIN_PASSWORD를 설정해야 합니다."
                )
            c.execute("INSERT INTO users(login_id,password_hash,name,role,is_active,must_change_password,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?)",
                      ("admin", hash_password(bootstrap_password), "초기 관리자", "ADMIN", 1, 1, now(), now()))


@app.on_event("startup")
def startup() -> None:
    init_db()


def row_dict(row: sqlite3.Row | None):
    return dict(row) if row else None


def audit(c: sqlite3.Connection, actor: int | None, action: str, target_type: str, target_id: int | str, before=None, after=None, ip: str | None = None):
    import json
    c.execute("INSERT INTO audit_logs(actor_user_id,action,target_type,target_id,before_data,after_data,ip_address,created_at) VALUES(?,?,?,?,?,?,?,?)",
              (actor, action, target_type, str(target_id), json.dumps(before, ensure_ascii=False) if before is not None else None,
               json.dumps(after, ensure_ascii=False) if after is not None else None, ip, now()))


SENSITIVE_AUDIT_KEYS = {
    "password", "password_hash", "temporary_password", "current_password", "new_password", "token", "storage_key",
}


def audit_value(value):
    """Decode audit JSON for administrators without ever returning sensitive values."""
    import json
    if value is None:
        return None
    try:
        decoded = json.loads(value)
    except (TypeError, ValueError):
        return None

    def redact(item):
        if isinstance(item, dict):
            return {key: "[REDACTED]" if key.lower() in SENSITIVE_AUDIT_KEYS else redact(child) for key, child in item.items()}
        if isinstance(item, list):
            return [redact(child) for child in item]
        return item

    return redact(decoded)


def audit_log_view(row: sqlite3.Row) -> dict:
    """Expose useful change detail while omitting the raw stored JSON columns."""
    out = dict(row)
    out["before"] = audit_value(out.pop("before_data"))
    out["after"] = audit_value(out.pop("after_data"))
    return out


def current_user(request: Request) -> sqlite3.Row:
    token = request.headers.get("Authorization", "").removeprefix("Bearer ").strip()
    session = SESSIONS.get(token)
    if not session or session[1] <= datetime.now(timezone.utc):
        raise HTTPException(status_code=401, detail="로그인이 필요합니다.")
    with db() as c:
        user = c.execute("SELECT * FROM users WHERE id=? AND is_active=1", (session[0],)).fetchone()
    if not user:
        SESSIONS.pop(token, None)
        raise HTTPException(status_code=401, detail="사용할 수 없는 계정입니다.")
    return user


def admin_required(user: sqlite3.Row = Depends(current_user)) -> sqlite3.Row:
    user = password_changed_user(user)
    if user["role"] != "ADMIN":
        raise HTTPException(status_code=403, detail="관리자 권한이 필요합니다.")
    return user


def password_changed_user(user: sqlite3.Row) -> sqlite3.Row:
    """Prevent a temporary password from being used for work data access."""
    if user["must_change_password"]:
        raise HTTPException(status_code=403, detail="최초 로그인 후 비밀번호를 변경해야 합니다.")
    return user


app.include_router(
    create_auth_router(
        AuthService(db, verify_password, password_hash_needs_upgrade, hash_password, now, SESSIONS, audit),
        current_user,
    )
)
app.include_router(
    create_users_router(
        UserService(db, verify_password, hash_password, now, SESSIONS, audit, sqlite3.IntegrityError),
        current_user,
        admin_required,
        password_changed_user,
    )
)


def accessible_task(c: sqlite3.Connection, task_id: int, user: sqlite3.Row) -> sqlite3.Row:
    task = c.execute("SELECT * FROM tasks WHERE id=? AND deleted_at IS NULL", (task_id,)).fetchone()
    if not task:
        raise HTTPException(status_code=404, detail="업무를 찾을 수 없습니다.")
    if user["role"] == "ADMIN" or task["created_by"] == user["id"]:
        return task
    assignee = c.execute("SELECT 1 FROM task_assignees WHERE task_id=? AND user_id=?", (task_id, user["id"])).fetchone()
    if not assignee:
        raise HTTPException(status_code=403, detail="해당 업무에 접근할 권한이 없습니다.")
    return task



@app.get("/")
def index():
    return FileResponse(STATIC_DIR / "index.html")


@app.get("/api/me/notification-preferences")
def notification_preferences(user: sqlite3.Row = Depends(current_user)):
    user = password_changed_user(user)
    with db() as c:
        return notification_preferences_view(c, user["id"])


@app.patch("/api/me/notification-preferences")
def patch_notification_preferences(
    payload: NotificationPreferencesIn, request: Request, user: sqlite3.Row = Depends(current_user)
):
    user = password_changed_user(user)
    with db() as c:
        before = notification_preferences_view(c, user["id"])
        c.execute(
            "UPDATE notification_preferences SET due_soon_enabled=?,overdue_enabled=?,updated_at=? WHERE user_id=?",
            (int(payload.due_soon_enabled), int(payload.overdue_enabled), now(), user["id"]),
        )
        after = notification_preferences_view(c, user["id"])
        # Disabling a local category also removes unread/read local entries from it;
        # no message is queued for an external service.
        if not payload.due_soon_enabled:
            c.execute("DELETE FROM notification_outbox WHERE user_id=? AND kind IN ('DUE_SOON','DUE_TODAY')", (user["id"],))
        if not payload.overdue_enabled:
            c.execute("DELETE FROM notification_outbox WHERE user_id=? AND kind='OVERDUE'", (user["id"],))
        if before != after:
            audit(c, user["id"], "NOTIFICATION_PREFERENCES_UPDATED", "USER", user["id"], before, after,
                  request.client.host if request.client else None)
        return after


@app.get("/api/notifications")
def notifications(
    unread_only: bool = False,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    user: sqlite3.Row = Depends(current_user),
):
    """Return only the caller's private local in-app due-date outbox entries."""
    user = password_changed_user(user)
    with db() as c:
        refresh_due_notifications(c, user)
        where = "WHERE n.user_id=?"
        params: list[int] = [user["id"]]
        if unread_only:
            where += " AND n.read_at IS NULL"
        total = c.execute(f"SELECT COUNT(*) FROM notification_outbox n {where}", params).fetchone()[0]
        rows = c.execute(
            "SELECT n.id,n.task_id,n.kind,n.due_date,n.message,n.read_at,n.created_at,t.description,t.status "
            "FROM notification_outbox n JOIN tasks t ON t.id=n.task_id " + where +
            " ORDER BY n.read_at IS NOT NULL,n.id DESC LIMIT ? OFFSET ?",
            [*params, page_size, (page - 1) * page_size],
        )
        return {
            "items": [dict(row) | {"is_read": row["read_at"] is not None} for row in rows],
            "pagination": {"page": page, "page_size": page_size, "total": total, "total_pages": (total + page_size - 1) // page_size},
            "delivery": "IN_APP_OUTBOX_ONLY",
        }


@app.post("/api/notifications/{notification_id}/read")
def read_notification(notification_id: int, user: sqlite3.Row = Depends(current_user)):
    user = password_changed_user(user)
    with db() as c:
        updated = c.execute(
            "UPDATE notification_outbox SET read_at=COALESCE(read_at,?) WHERE id=? AND user_id=?",
            (now(), notification_id, user["id"]),
        )
        if not updated.rowcount:
            raise HTTPException(status_code=404, detail="알림을 찾을 수 없습니다.")
        return {"ok": True}


def organization_view(row: sqlite3.Row) -> dict:
    out = dict(row)
    out["is_active"] = bool(out["is_active"])
    return out


def contact_view(c: sqlite3.Connection, row: sqlite3.Row) -> dict:
    out = dict(row)
    out["is_active"] = bool(out["is_active"])
    out["organization"] = c.execute("SELECT name FROM organizations WHERE id=?", (row["organization_id"],)).fetchone()["name"]
    return out


def contact_from_input(c: sqlite3.Connection, payload: ContactIn, actor_id: int | None = None, ip: str | None = None) -> int:
    """Find or create reusable master records, recording only actual master changes."""
    org = c.execute("SELECT * FROM organizations WHERE name=?", (payload.organization,)).fetchone()
    if not org:
        timestamp = now()
        cur = c.execute("INSERT INTO organizations(name,is_active,created_at,updated_at) VALUES(?,1,?,?)", (payload.organization, timestamp, timestamp))
        org_id = cur.lastrowid
        audit(c, actor_id, "ORGANIZATION_CREATED", "ORGANIZATION", org_id,
              after={"name": payload.organization, "is_active": True}, ip=ip)
    else:
        org_id = org["id"]
    contact = c.execute("SELECT * FROM request_contacts WHERE name=? AND phone=? AND organization_id=?", (payload.name, payload.phone, org_id)).fetchone()
    if contact:
        return contact["id"]
    timestamp = now()
    contact_id = c.execute(
        "INSERT INTO request_contacts(name,phone,organization_id,is_active,created_at,updated_at) VALUES(?,?,?,1,?,?)",
        (payload.name, payload.phone, org_id, timestamp, timestamp),
    ).lastrowid
    audit(c, actor_id, "REQUEST_CONTACT_CREATED", "REQUEST_CONTACT", contact_id,
          after={"name": payload.name, "phone": payload.phone, "organization_id": org_id, "is_active": True}, ip=ip)
    return contact_id

def image_view(row: sqlite3.Row) -> dict:
    """Metadata deliberately excludes the private storage key."""
    return {
        "id": row["id"], "original_filename": row["original_filename"], "mime_type": row["mime_type"],
        "size_bytes": row["size_bytes"], "sort_order": row["sort_order"], "created_at": row["created_at"],
    }


def attachment_view(row: sqlite3.Row) -> dict:
    """Public attachment metadata; private paths and checksums never leave the API."""
    return {
        "id": row["id"], "original_filename": row["original_filename"], "mime_type": row["mime_type"],
        "attachment_type": row["attachment_type"], "size_bytes": row["size_bytes"], "created_at": row["created_at"],
    }


def private_upload_path(storage_key: str, missing_detail: str = "첨부 파일을 찾을 수 없습니다.") -> Path:
    root = UPLOAD_DIR.resolve()
    candidate = (root / storage_key).resolve()
    if root != candidate and root not in candidate.parents:
        raise HTTPException(status_code=404, detail=missing_detail)
    return candidate


def private_image_path(storage_key: str) -> Path:
    return private_upload_path(storage_key, "이미지를 찾을 수 없습니다.")


def validate_image_upload(upload: UploadFile, content: bytes) -> str:
    claimed_type = (upload.content_type or "").lower()
    signature_check = IMAGE_SIGNATURES.get(claimed_type)
    if not signature_check or not signature_check(content):
        raise HTTPException(
            status_code=422,
            detail="이미지는 JPG, PNG, WEBP 파일만 허용하며 파일 형식과 실제 시그니처가 일치해야 합니다.",
        )
    return claimed_type


def rotate_private_image(content: bytes, mime_type: str, direction: str) -> bytes:
    """Decode, rotate, and re-encode a private image within strict local limits.

    The source was MIME/signature checked at upload time, but decoding is still
    defensive because a compressed image can expand dramatically. Metadata is not
    copied into the derivative, avoiding accidental retention of EXIF location data.
    """
    formats = {"image/jpeg": "JPEG", "image/png": "PNG", "image/webp": "WEBP"}
    image_format = formats.get(mime_type)
    if not image_format:
        raise HTTPException(status_code=422, detail="편집할 수 없는 이미지 형식입니다.")
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(BytesIO(content)) as source:
                source.load()
                if source.width * source.height > MAX_IMAGE_PIXELS:
                    raise Image.DecompressionBombError("image dimensions exceed the editing limit")
                operation = Image.Transpose.ROTATE_270 if direction == "CLOCKWISE" else Image.Transpose.ROTATE_90
                edited = source.transpose(operation)
                if image_format == "JPEG" and edited.mode not in ("RGB", "L"):
                    edited = edited.convert("RGB")
                elif image_format == "WEBP" and edited.mode not in ("RGB", "RGBA"):
                    edited = edited.convert("RGBA" if "transparency" in source.info else "RGB")
                result = BytesIO()
                save_options = {"format": image_format}
                if image_format == "JPEG":
                    save_options.update({"quality": 90, "optimize": True})
                elif image_format == "WEBP":
                    save_options.update({"quality": 90, "method": 4})
                else:
                    save_options["optimize"] = True
                edited.save(result, **save_options)
                output = result.getvalue()
    except (Image.DecompressionBombError, Image.DecompressionBombWarning, UnidentifiedImageError, OSError, ValueError):
        raise HTTPException(status_code=422, detail="안전하게 편집할 수 없는 이미지입니다.")
    if not output or len(output) > MAX_IMAGE_BYTES or not IMAGE_SIGNATURES[mime_type](output):
        raise HTTPException(status_code=422, detail="편집 결과가 허용된 이미지 형식 또는 크기를 충족하지 않습니다.")
    return output


def valid_office_zip(content: bytes, required_member: str) -> bool:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            names = set(archive.namelist())
            return "[Content_Types].xml" in names and required_member in names
    except (OSError, zipfile.BadZipFile):
        return False


def validate_attachment_upload(upload: UploadFile, content: bytes) -> tuple[str, str, str]:
    claimed_type = (upload.content_type or "").lower()
    rule = ATTACHMENT_RULES.get(claimed_type)
    if not rule:
        raise HTTPException(status_code=422, detail="문서는 PDF, TXT, DOCX, XLSX만, 동영상은 MP4, WEBM만 허용합니다.")
    attachment_type, suffix, _ = rule
    valid = (
        (claimed_type == "application/pdf" and content.startswith(b"%PDF-"))
        or (claimed_type == "text/plain" and b"\x00" not in content and _is_utf8_text(content))
        or (claimed_type.endswith("wordprocessingml.document") and valid_office_zip(content, "word/document.xml"))
        or (claimed_type.endswith("spreadsheetml.sheet") and valid_office_zip(content, "xl/workbook.xml"))
        or (claimed_type == "video/mp4" and len(content) >= 12 and content[4:8] == b"ftyp")
        or (claimed_type == "video/webm" and content.startswith(b"\x1aE\xdf\xa3"))
    )
    if not valid:
        raise HTTPException(status_code=422, detail="첨부 파일의 MIME 형식과 실제 파일 구조가 일치하지 않습니다.")
    return attachment_type, suffix, claimed_type


def _is_utf8_text(content: bytes) -> bool:
    try:
        content.decode("utf-8")
        return True
    except UnicodeDecodeError:
        return False


def task_view(c: sqlite3.Connection, task: sqlite3.Row) -> dict:
    out = dict(task)
    out["created_by_name"] = c.execute("SELECT name FROM users WHERE id=?", (task["created_by"],)).fetchone()["name"]
    out["assignees"] = [dict(row) for row in c.execute("SELECT u.id,u.name,u.login_id FROM task_assignees ta JOIN users u ON u.id=ta.user_id WHERE ta.task_id=? ORDER BY u.name", (task["id"],))]
    out["contact"] = row_dict(c.execute("SELECT rc.id,rc.name,rc.phone,o.name organization FROM request_contacts rc JOIN organizations o ON o.id=rc.organization_id WHERE rc.id=?", (task["requester_contact_id"],)).fetchone())
    out["images"] = [image_view(row) for row in c.execute(
        "SELECT * FROM task_images WHERE task_id=? ORDER BY sort_order,id", (task["id"],)
    )]
    out["attachments"] = [attachment_view(row) for row in c.execute(
        "SELECT * FROM task_attachments WHERE task_id=? ORDER BY created_at,id", (task["id"],)
    )]
    return out


def task_changes(before: dict, after: dict) -> list[str]:
    labels = {
        "status": "상태", "description": "업무 설명", "received_date": "접수일",
        "due_date": "마감일", "completed_date": "처리일", "contact": "요청 정보", "assignees": "처리자",
    }
    return [labels[key] for key in labels if before.get(key) != after.get(key)]


def task_history(c: sqlite3.Connection, task_id: int) -> list[dict]:
    import json
    rows = c.execute(
        "SELECT a.id,a.action,a.before_data,a.after_data,a.created_at,u.name actor_name "
        "FROM audit_logs a LEFT JOIN users u ON u.id=a.actor_user_id "
        "WHERE a.target_type='TASK' AND a.target_id=? ORDER BY a.id DESC", (str(task_id),)
    )
    history = []
    for row in rows:
        item = dict(row)
        item["before"] = json.loads(item.pop("before_data")) if item["before_data"] else None
        item["after"] = json.loads(item.pop("after_data")) if item["after_data"] else None
        item["changed_fields"] = task_changes(item["before"], item["after"]) if item["before"] and item["after"] else []
        history.append(item)
    return history


def task_query_sql(user: sqlite3.Row) -> tuple[str, list]:
    sql = "SELECT t.* FROM tasks t WHERE t.deleted_at IS NULL"
    params=[]
    if user["role"] != "ADMIN":
        sql += " AND (t.created_by=? OR EXISTS(SELECT 1 FROM task_assignees ta WHERE ta.task_id=t.id AND ta.user_id=?))"
        params=[user["id"], user["id"]]
    return sql, params


def excel_text(value) -> str:
    """Return spreadsheet-safe text without allowing a task value to become a formula."""
    text = "" if value is None else str(value)
    return f"'{text}" if text.startswith(("=", "+", "-", "@")) else text


def status_label(task_status: str) -> str:
    return {"RECEIVED": "접수", "IN_PROGRESS": "진행중", "COMPLETED": "처리"}[task_status]


def task_export_workbook(c: sqlite3.Connection, rows, exported_by: sqlite3.Row) -> bytes:
    """Build a local XLSX workbook containing only already-authorized task rows."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "업무 목록"
    headers = [
        "업무 ID", "상태", "업무 설명", "요청 담당자", "연락처", "소속 기관", "처리자",
        "접수일", "마감일", "처리일", "생성자", "생성 시각", "수정 시각",
    ]
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="2463A6")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for task in rows:
        view = task_view(c, task)
        sheet.append([
            view["id"], status_label(view["status"]), excel_text(view["description"]),
            excel_text(view["contact"]["name"]), excel_text(view["contact"]["phone"]),
            excel_text(view["contact"]["organization"]),
            excel_text(", ".join(assignee["name"] for assignee in view["assignees"])),
            view["received_date"], view["due_date"] or "", view["completed_date"] or "",
            excel_text(view["created_by_name"]), view["created_at"], view["updated_at"],
        ])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [10, 12, 48, 18, 18, 24, 28, 14, 14, 14, 18, 28, 28]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    guide = workbook.create_sheet("내보내기 정보")
    guide.append(["항목", "값"])
    guide.append(["내보낸 시각(UTC)", now()])
    guide.append(["내보낸 사용자", excel_text(exported_by["name"])])
    guide.append(["업무 건수", len(rows)])
    guide.append(["보안", "이 파일은 로그인 사용자의 서버 권한 범위와 적용된 필터만 포함합니다."])
    for cell in guide[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["B"].width = 86
    for row in guide.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def notification_preferences_view(c: sqlite3.Connection, user_id: int) -> dict:
    """Return (and lazily create) a user's safe, local-only reminder settings."""
    c.execute(
        "INSERT INTO notification_preferences(user_id,due_soon_enabled,overdue_enabled,updated_at) "
        "VALUES(?,1,1,?) ON CONFLICT(user_id) DO NOTHING",
        (user_id, now()),
    )
    row = c.execute("SELECT * FROM notification_preferences WHERE user_id=?", (user_id,)).fetchone()
    return {
        "due_soon_enabled": bool(row["due_soon_enabled"]),
        "overdue_enabled": bool(row["overdue_enabled"]),
        "updated_at": row["updated_at"],
        "delivery": "IN_APP_OUTBOX_ONLY",
    }


def refresh_due_notifications(c: sqlite3.Connection, user: sqlite3.Row) -> None:
    """Populate a user's private in-app due-date outbox without external delivery."""
    preferences = notification_preferences_view(c, user["id"])
    if not preferences["due_soon_enabled"]:
        c.execute("DELETE FROM notification_outbox WHERE user_id=? AND kind IN ('DUE_SOON','DUE_TODAY')", (user["id"],))
    if not preferences["overdue_enabled"]:
        c.execute("DELETE FROM notification_outbox WHERE user_id=? AND kind='OVERDUE'", (user["id"],))
    task_sql, params = task_query_sql(user)
    task_sql += " AND t.status!='COMPLETED' AND t.due_date IS NOT NULL"
    today = today_in_seoul()
    for task in c.execute(task_sql, params):
        due = date.fromisoformat(task["due_date"])
        if due < today:
            kind = "OVERDUE"
            enabled = preferences["overdue_enabled"]
            message = f"마감일이 지났습니다: {task['description'][:80]}"
        elif due == today:
            kind = "DUE_TODAY"
            enabled = preferences["due_soon_enabled"]
            message = f"오늘 마감: {task['description'][:80]}"
        elif due <= today + timedelta(days=3):
            kind = "DUE_SOON"
            enabled = preferences["due_soon_enabled"]
            message = f"마감 임박 ({due.isoformat()}): {task['description'][:80]}"
        else:
            continue
        if enabled:
            c.execute(
                "INSERT INTO notification_outbox(user_id,task_id,kind,due_date,message,created_at) "
                "VALUES(?,?,?,?,?,?) ON CONFLICT(user_id,task_id,kind,due_date) DO NOTHING",
                (user["id"], task["id"], kind, str(due), message, now()),
            )


def apply_statistics_filters(
    sql: str,
    params: list,
    user: sqlite3.Row,
    received_from: date | None,
    received_to: date | None,
    assignee_id: int | None,
) -> tuple[str, list, int | None]:
    """Apply report scope without weakening the existing task visibility policy."""
    if received_from and received_to and received_from > received_to:
        raise HTTPException(status_code=422, detail="접수일 시작일은 종료일보다 늦을 수 없습니다.")
    if user["role"] != "ADMIN" and assignee_id is not None and assignee_id != user["id"]:
        raise HTTPException(status_code=403, detail="일반 이용자의 처리자 필터는 본인으로 고정됩니다.")
    if received_from:
        sql += " AND t.received_date>=?"
        params.append(str(received_from))
    if received_to:
        sql += " AND t.received_date<=?"
        params.append(str(received_to))
    if assignee_id is not None:
        sql += " AND EXISTS(SELECT 1 FROM task_assignees ta_filter WHERE ta_filter.task_id=t.id AND ta_filter.user_id=?)"
        params.append(assignee_id)
    # A member may have creator-only tasks in the status aggregate. Their per-user
    # row remains self-only, so it cannot reveal shared processors' identities.
    user_scope = assignee_id if assignee_id is not None else (user["id"] if user["role"] != "ADMIN" else None)
    return sql, params, user_scope


@app.get("/api/statistics")
def statistics(
    received_from: date | None = None,
    received_to: date | None = None,
    assignee_id: int | None = Query(None, gt=0),
    user: sqlite3.Row = Depends(current_user),
):
    """Return status and processor aggregates for only the caller's task scope."""
    user = password_changed_user(user)
    with db() as c:
        sql, params = task_query_sql(user)
        sql, params, user_scope = apply_statistics_filters(
            sql, params, user, received_from, received_to, assignee_id
        )
        by_status = {name: 0 for name in ("RECEIVED", "IN_PROGRESS", "COMPLETED")}
        for row in c.execute(f"SELECT visible.status,COUNT(*) count FROM ({sql}) visible GROUP BY visible.status", params):
            by_status[row["status"]] = row["count"]
        total = sum(by_status.values())

        user_sql = (
            "SELECT u.id,u.name,u.login_id,u.is_active,COUNT(*) total,"
            "SUM(CASE WHEN visible.status='RECEIVED' THEN 1 ELSE 0 END) RECEIVED,"
            "SUM(CASE WHEN visible.status='IN_PROGRESS' THEN 1 ELSE 0 END) IN_PROGRESS,"
            "SUM(CASE WHEN visible.status='COMPLETED' THEN 1 ELSE 0 END) COMPLETED "
            "FROM task_assignees ta JOIN users u ON u.id=ta.user_id "
            f"JOIN ({sql}) visible ON visible.id=ta.task_id"
        )
        user_params = list(params)
        if user_scope is not None:
            user_sql += " WHERE ta.user_id=?"
            user_params.append(user_scope)
        user_sql += " GROUP BY u.id,u.name,u.login_id,u.is_active ORDER BY total DESC,u.name,u.id"
        by_user = []
        for row in c.execute(user_sql, user_params):
            item = dict(row)
            item["is_active"] = bool(item["is_active"])
            item["by_status"] = {name: item.pop(name) for name in ("RECEIVED", "IN_PROGRESS", "COMPLETED")}
            by_user.append(item)
        return {
            "scope": "ALL_TASKS" if user["role"] == "ADMIN" and assignee_id is None else "MY_AUTHORIZED_TASKS",
            "period": {"received_from": str(received_from) if received_from else None, "received_to": str(received_to) if received_to else None},
            "assignee_id": assignee_id,
            "total": total,
            "by_status": by_status,
            "by_user": by_user,
            "user_totals_note": "공동 처리 업무는 각 처리자에게 1건씩 집계됩니다.",
        }


@app.get("/api/tasks/export.xlsx")
def export_tasks_xlsx(
    request: Request,
    status: Literal["RECEIVED", "IN_PROGRESS", "COMPLETED"] | None = None,
    q: str | None = Query(None, max_length=200),
    received_from: date | None = None,
    received_to: date | None = None,
    due_from: date | None = None,
    due_to: date | None = None,
    assignee_id: int | None = Query(None, gt=0),
    user: sqlite3.Row = Depends(current_user),
):
    """Create a download-only XLSX for the same filtered server-side task scope."""
    user = password_changed_user(user)
    if received_from and received_to and received_from > received_to:
        raise HTTPException(status_code=422, detail="접수일 시작일은 종료일보다 늦을 수 없습니다.")
    if due_from and due_to and due_from > due_to:
        raise HTTPException(status_code=422, detail="마감일 시작일은 종료일보다 늦을 수 없습니다.")
    if user["role"] != "ADMIN" and assignee_id is not None and assignee_id != user["id"]:
        raise HTTPException(status_code=403, detail="일반 이용자의 처리자 필터는 본인으로 고정됩니다.")
    with db() as c:
        sql, params = task_query_sql(user)
        if status:
            sql += " AND t.status=?"
            params.append(status)
        if received_from:
            sql += " AND t.received_date>=?"
            params.append(str(received_from))
        if received_to:
            sql += " AND t.received_date<=?"
            params.append(str(received_to))
        if due_from:
            sql += " AND t.due_date>=?"
            params.append(str(due_from))
        if due_to:
            sql += " AND t.due_date<=?"
            params.append(str(due_to))
        keyword = (q or "").strip()
        if keyword:
            sql += " AND (t.description LIKE ? OR EXISTS(SELECT 1 FROM request_contacts rc JOIN organizations o ON o.id=rc.organization_id WHERE rc.id=t.requester_contact_id AND (rc.name LIKE ? OR o.name LIKE ?)))"
            params += [f"%{keyword}%"] * 3
        if assignee_id is not None:
            sql += " AND EXISTS(SELECT 1 FROM task_assignees ta_filter WHERE ta_filter.task_id=t.id AND ta_filter.user_id=?)"
            params.append(assignee_id)
        rows = list(c.execute(
            sql + " ORDER BY CASE t.status WHEN 'RECEIVED' THEN 1 WHEN 'IN_PROGRESS' THEN 2 ELSE 3 END,t.received_date DESC,t.id DESC LIMIT 10001",
            params,
        ))
        if len(rows) > 10_000:
            raise HTTPException(status_code=422, detail="엑셀 내보내기는 한 번에 최대 10,000건까지 가능합니다. 필터를 좁혀 다시 시도하세요.")
        document = task_export_workbook(c, rows, user)
        audit(c, user["id"], "TASKS_EXPORTED", "TASK_EXPORT", "FILTERED_SCOPE", after={
            "count": len(rows), "status": status, "received_from": str(received_from) if received_from else None,
            "received_to": str(received_to) if received_to else None, "due_from": str(due_from) if due_from else None,
            "due_to": str(due_to) if due_to else None, "assignee_id": assignee_id, "has_keyword": bool(keyword),
        }, ip=request.client.host if request.client else None)
    return Response(
        content=document,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": "attachment; filename=shared-work-calendar-tasks.xlsx",
            "Cache-Control": "private, no-store",
            "X-Content-Type-Options": "nosniff",
        },
    )


@app.post("/api/tasks/{task_id}/attachments", status_code=201)
async def upload_task_attachments(
    task_id: int, request: Request, attachments: list[UploadFile] = File(...), user: sqlite3.Row = Depends(current_user)
):
    """Store validated private documents/videos; content is served as download only."""
    user = password_changed_user(user)
    if not attachments or len(attachments) > MAX_ATTACHMENT_COUNT:
        raise HTTPException(status_code=422, detail="한 번에 문서/동영상 첨부는 1개 이상 10개 이하만 업로드할 수 있습니다.")
    # Reject a direct URL/API authorization bypass before inspecting private bytes.
    with db() as c:
        accessible_task(c, task_id, user)
    prepared: list[tuple[UploadFile, bytes, str, str, str]] = []
    created_paths: list[Path] = []
    try:
        for upload in attachments:
            rule = ATTACHMENT_RULES.get((upload.content_type or "").lower())
            if not rule:
                raise HTTPException(status_code=422, detail="문서는 PDF, TXT, DOCX, XLSX만, 동영상은 MP4, WEBM만 허용합니다.")
            content = await upload.read(rule[2] + 1)
            if not content or len(content) > rule[2]:
                limit_mb = rule[2] // (1024 * 1024)
                raise HTTPException(status_code=422, detail=f"{('문서' if rule[0] == 'DOCUMENT' else '동영상')} 파일당 최대 크기는 {limit_mb}MB입니다.")
            attachment_type, suffix, mime_type = validate_attachment_upload(upload, content)
            prepared.append((upload, content, attachment_type, suffix, mime_type))
        with db() as c:
            accessible_task(c, task_id, user)
            existing_count = c.execute("SELECT COUNT(*) FROM task_attachments WHERE task_id=?", (task_id,)).fetchone()[0]
            if existing_count + len(prepared) > MAX_ATTACHMENT_COUNT:
                raise HTTPException(status_code=422, detail="업무당 문서/동영상 첨부는 최대 10개입니다.")
            UPLOAD_DIR.mkdir(mode=0o700, parents=True, exist_ok=True)
            uploaded = []
            for upload, content, attachment_type, suffix, mime_type in prepared:
                storage_key = f"attachments/{task_id}/{uuid.uuid4().hex}{suffix}"
                destination = private_upload_path(storage_key)
                destination.parent.mkdir(mode=0o700, parents=True, exist_ok=True)
                destination.write_bytes(content)
                try:
                    os.chmod(destination, 0o600)
                except OSError:
                    pass
                created_paths.append(destination)
                filename = Path(upload.filename or f"attachment{suffix}").name[-255:] or f"attachment{suffix}"
                attachment_id = c.execute(
                    "INSERT INTO task_attachments(task_id,storage_key,original_filename,mime_type,attachment_type,size_bytes,sha256,uploaded_by,created_at) VALUES(?,?,?,?,?,?,?,?,?)",
                    (task_id, storage_key, filename, mime_type, attachment_type, len(content), hashlib.sha256(content).hexdigest(), user["id"], now()),
                ).lastrowid
                uploaded.append(attachment_view(c.execute("SELECT * FROM task_attachments WHERE id=?", (attachment_id,)).fetchone()))
            audit(c, user["id"], "TASK_ATTACHMENT_UPLOADED", "TASK", task_id, after={"attachments": uploaded},
                  ip=request.client.host if request.client else None)
            return {"attachments": uploaded}
    except sqlite3.IntegrityError as exc:
        for path in created_paths:
            path.unlink(missing_ok=True)
        raise HTTPException(status_code=422, detail="업무당 문서/동영상 첨부는 최대 10개입니다.") from exc
    except Exception:
        for path in created_paths:
            path.unlink(missing_ok=True)
        raise
    finally:
        for upload in attachments:
            await upload.close()


@app.delete("/api/tasks/{task_id}/attachments/{attachment_id}")
def delete_task_attachment(task_id: int, attachment_id: int, request: Request, user: sqlite3.Row = Depends(current_user)):
    user = password_changed_user(user)
    with db() as c:
        accessible_task(c, task_id, user)
        attachment = c.execute("SELECT * FROM task_attachments WHERE id=? AND task_id=?", (attachment_id, task_id)).fetchone()
        if not attachment:
            raise HTTPException(status_code=404, detail="첨부 파일을 찾을 수 없습니다.")
        result = attachment_view(attachment)
        path = private_upload_path(attachment["storage_key"])
        c.execute("DELETE FROM task_attachments WHERE id=?", (attachment_id,))
        audit(c, user["id"], "TASK_ATTACHMENT_DELETED", "TASK", task_id, before={"attachment": result},
              ip=request.client.host if request.client else None)
    path.unlink(missing_ok=True)
    return {"ok": True}


@app.get("/api/completed-calendar")
def completed_calendar(month: str, user: sqlite3.Row=Depends(current_user)):
    """Return only caller-authorized completed tasks for a strictly validated month."""
    user = password_changed_user(user)
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", month):
        raise HTTPException(status_code=422, detail="month는 유효한 YYYY-MM 형식이어야 합니다.")
    year, month_number = map(int, month.split("-"))
    first_day = date(year, month_number, 1)
    next_month = date(year + 1, 1, 1) if month_number == 12 else date(year, month_number + 1, 1)
    with db() as c:
        sql, params = task_query_sql(user)
        sql += " AND t.status='COMPLETED' AND t.completed_date>=? AND t.completed_date<? ORDER BY t.completed_date,t.id"
        params.extend([str(first_day), str(next_month)])
        return [task_view(c, row) for row in c.execute(sql, params)]


@app.get("/api/organizations")
def organizations(q: str="", user: sqlite3.Row=Depends(current_user)):
    user = password_changed_user(user)
    with db() as c:
        rows = c.execute(
            "SELECT * FROM organizations WHERE is_active=1 AND name LIKE ? ORDER BY name LIMIT 20",
            (f"%{q.strip()}%",),
        )
        return [organization_view(row) for row in rows]


@app.post("/api/organizations", status_code=201)
def create_organization(payload: OrganizationIn, request: Request, user: sqlite3.Row=Depends(current_user)):
    user = password_changed_user(user)
    with db() as c:
        existing = c.execute("SELECT * FROM organizations WHERE name=?", (payload.name,)).fetchone()
        if existing:
            return organization_view(existing)
        timestamp = now()
        org_id = c.execute(
            "INSERT INTO organizations(name,is_active,created_at,updated_at) VALUES(?,1,?,?)",
            (payload.name, timestamp, timestamp),
        ).lastrowid
        row = c.execute("SELECT * FROM organizations WHERE id=?", (org_id,)).fetchone()
        result = organization_view(row)
        audit(c, user["id"], "ORGANIZATION_CREATED", "ORGANIZATION", org_id, after=result,
              ip=request.client.host if request.client else None)
        return result


@app.get("/api/request-contacts")
def contacts(q: str="", user: sqlite3.Row=Depends(current_user)):
    user = password_changed_user(user)
    with db() as c:
        rows = c.execute(
            "SELECT rc.* FROM request_contacts rc JOIN organizations o ON o.id=rc.organization_id "
            "WHERE rc.is_active=1 AND o.is_active=1 AND (rc.name LIKE ? OR o.name LIKE ?) "
            "ORDER BY rc.name, o.name LIMIT 20", (f"%{q.strip()}%", f"%{q.strip()}%"),
        )
        return [contact_view(c, row) for row in rows]


@app.post("/api/request-contacts", status_code=201)
def create_contact(payload: ContactIn, request: Request, user: sqlite3.Row=Depends(current_user)):
    user = password_changed_user(user)
    with db() as c:
        contact_id = contact_from_input(c, payload, user["id"], request.client.host if request.client else None)
        return contact_view(c, c.execute("SELECT * FROM request_contacts WHERE id=?", (contact_id,)).fetchone())


@app.get("/api/admin/organizations")
def admin_organizations(q: str="", _: sqlite3.Row=Depends(admin_required)):
    with db() as c:
        return [organization_view(row) for row in c.execute(
            "SELECT * FROM organizations WHERE name LIKE ? ORDER BY is_active DESC,name LIMIT 100", (f"%{q.strip()}%",)
        )]


@app.patch("/api/admin/organizations/{organization_id}")
def patch_organization(organization_id: int, payload: OrganizationPatch, request: Request, user: sqlite3.Row=Depends(admin_required)):
    with db() as c:
        old = c.execute("SELECT * FROM organizations WHERE id=?", (organization_id,)).fetchone()
        if not old:
            raise HTTPException(status_code=404, detail="기관을 찾을 수 없습니다.")
        new_name = payload.name if payload.name is not None else old["name"]
        new_active = int(old["is_active"] if payload.is_active is None else payload.is_active)
        conflict = c.execute("SELECT id FROM organizations WHERE name=? AND id!=?", (new_name, organization_id)).fetchone()
        if conflict:
            raise HTTPException(status_code=409, detail="같은 기관명이 이미 존재합니다.")
        before = organization_view(old)
        c.execute("UPDATE organizations SET name=?,is_active=?,updated_at=? WHERE id=?", (new_name, new_active, now(), organization_id))
        updated = organization_view(c.execute("SELECT * FROM organizations WHERE id=?", (organization_id,)).fetchone())
        if before != updated:
            audit(c, user["id"], "ORGANIZATION_UPDATED", "ORGANIZATION", organization_id, before, updated,
                  request.client.host if request.client else None)
        return updated


@app.get("/api/admin/request-contacts")
def admin_contacts(q: str="", _: sqlite3.Row=Depends(admin_required)):
    with db() as c:
        rows = c.execute(
            "SELECT rc.* FROM request_contacts rc JOIN organizations o ON o.id=rc.organization_id "
            "WHERE rc.name LIKE ? OR rc.phone LIKE ? OR o.name LIKE ? "
            "ORDER BY rc.is_active DESC,rc.name,o.name LIMIT 100", (f"%{q.strip()}%",) * 3,
        )
        return [contact_view(c, row) for row in rows]


@app.patch("/api/admin/request-contacts/{contact_id}")
def patch_contact(contact_id: int, payload: ContactPatch, request: Request, user: sqlite3.Row=Depends(admin_required)):
    with db() as c:
        old = c.execute("SELECT * FROM request_contacts WHERE id=?", (contact_id,)).fetchone()
        if not old:
            raise HTTPException(status_code=404, detail="요청 담당자를 찾을 수 없습니다.")
        name = payload.name if payload.name is not None else old["name"]
        phone = payload.phone if payload.phone is not None else old["phone"]
        org_id = payload.organization_id if payload.organization_id is not None else old["organization_id"]
        is_active = int(old["is_active"] if payload.is_active is None else payload.is_active)
        if not c.execute("SELECT 1 FROM organizations WHERE id=?", (org_id,)).fetchone():
            raise HTTPException(status_code=422, detail="존재하지 않는 기관입니다.")
        duplicate = c.execute("SELECT id FROM request_contacts WHERE name=? AND phone=? AND organization_id=? AND id!=?", (name, phone, org_id, contact_id)).fetchone()
        if duplicate:
            raise HTTPException(status_code=409, detail="같은 이름, 연락처, 기관의 요청 담당자가 이미 존재합니다.")
        before = contact_view(c, old)
        c.execute("UPDATE request_contacts SET name=?,phone=?,organization_id=?,is_active=?,updated_at=? WHERE id=?", (name, phone, org_id, is_active, now(), contact_id))
        updated = contact_view(c, c.execute("SELECT * FROM request_contacts WHERE id=?", (contact_id,)).fetchone())
        if before != updated:
            audit(c, user["id"], "REQUEST_CONTACT_UPDATED", "REQUEST_CONTACT", contact_id, before, updated,
                  request.client.host if request.client else None)
        return updated


@app.get("/api/admin/audit-logs")
def audit_logs(
    action: str | None = Query(None, min_length=1, max_length=50, pattern=r"^[A-Z][A-Z0-9_]*$"),
    actor_user_id: int | None = Query(None, gt=0),
    created_from: date | None = None,
    created_to: date | None = None,
    _: sqlite3.Row = Depends(admin_required),
):
    """Administrator-only audit trail filtered by event, actor, and inclusive UTC date range."""
    if created_from and created_to and created_from > created_to:
        raise HTTPException(status_code=422, detail="시작일은 종료일보다 늦을 수 없습니다.")
    sql = "SELECT a.*,u.name actor_name FROM audit_logs a LEFT JOIN users u ON u.id=a.actor_user_id WHERE 1=1"
    params: list[str | int] = []
    if action:
        sql += " AND a.action=?"
        params.append(action)
    if actor_user_id is not None:
        sql += " AND a.actor_user_id=?"
        params.append(actor_user_id)
    if created_from:
        sql += " AND a.created_at>=?"
        params.append(str(created_from))
    if created_to:
        sql += " AND a.created_at<?"
        params.append(str(created_to + timedelta(days=1)))
    sql += " ORDER BY a.id DESC LIMIT 100"
    with db() as c:
        return [audit_log_view(row) for row in c.execute(sql, params)]


# Register after the existing static task-export route so its literal path retains
# precedence over the legacy ``/api/tasks/{task_id}`` parameter route.
app.include_router(
    create_task_router(
        TaskService(
            db,
            task_query_sql,
            task_view,
            accessible_task,
            contact_from_input,
            lambda: today_in_seoul(),
            lambda: now(),
            audit,
            task_changes,
            task_history,
            private_image_path,
            private_upload_path,
            image_view,
            rotate_private_image,
            IMAGE_SUFFIXES,
            UPLOAD_DIR,
            MAX_IMAGE_COUNT,
            MAX_IMAGE_BYTES,
            validate_image_upload,
            sqlite3.IntegrityError,
        ),
        current_user,
        password_changed_user,
        admin_required,
    )
)
