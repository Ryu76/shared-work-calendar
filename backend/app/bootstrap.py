from __future__ import annotations

import hashlib
import hmac
import os
import secrets
import uuid
import warnings
from io import BytesIO
from contextlib import contextmanager
from datetime import date, datetime, timezone
from pathlib import Path

from zoneinfo import ZoneInfo

from argon2 import PasswordHasher
from argon2.exceptions import InvalidHashError, VerificationError
from fastapi import Depends, HTTPException, Request, status
from fastapi.responses import FileResponse, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, UnidentifiedImageError

# Persistence is deliberately isolated: database.py enforces SQLCipher encryption
# and gets its 256-bit key from the user's Windows DPAPI-protected key file.
try:  # Imported as ``app.main`` with backend/ on sys.path.
    from core.config import load_settings
    from core.logging import configure_logging
    from app.factory import create_application
    from database import DB_PATH, connect, sqlite3
    from schemas.requests import TaskIn, TaskUpdateIn
    from routers.auth import create_router as create_auth_router
    from routers.dashboard import create_router as create_dashboard_router
    from routers.master import create_router as create_master_router
    from routers.notifications import create_router as create_notification_router
    from routers.tasks import create_router as create_task_router
    from routers.users import create_router as create_users_router
    from repositories.auth_repository import AuthRepository
    from repositories.audit_repository import AuditRepository
    from repositories.bootstrap_repository import BootstrapRepository
    from repositories.dashboard_repository import DashboardRepository
    from repositories.master_repository import MasterRepository
    from repositories.notification_repository import NotificationRepository
    from repositories.task_repository import TaskRepository
    from repositories.user_repository import UserRepository
    from services.auth_service import AuthService
    from services.dashboard_service import DashboardService
    from services.master_service import MasterService
    from services.notification_service import NotificationService
    from services.task_service import TaskService
    from services.task_projection_service import TaskProjectionService
    from services.user_service import UserService
    from validators.uploads import (
        ATTACHMENT_RULES,
        IMAGE_SIGNATURES,
        IMAGE_SUFFIXES,
        MAX_IMAGE_BYTES,
        validate_attachment_upload,
        validate_image_upload,
    )
except ModuleNotFoundError:  # Imported as ``backend.app.main`` from repository root.
    from backend.core.config import load_settings
    from backend.core.logging import configure_logging
    from backend.app.factory import create_application
    from backend.database import DB_PATH, connect, sqlite3
    from backend.schemas.requests import TaskIn, TaskUpdateIn
    from backend.routers.auth import create_router as create_auth_router
    from backend.routers.dashboard import create_router as create_dashboard_router
    from backend.routers.master import create_router as create_master_router
    from backend.routers.notifications import create_router as create_notification_router
    from backend.routers.tasks import create_router as create_task_router
    from backend.routers.users import create_router as create_users_router
    from backend.repositories.auth_repository import AuthRepository
    from backend.repositories.audit_repository import AuditRepository
    from backend.repositories.bootstrap_repository import BootstrapRepository
    from backend.repositories.dashboard_repository import DashboardRepository
    from backend.repositories.master_repository import MasterRepository
    from backend.repositories.notification_repository import NotificationRepository
    from backend.repositories.task_repository import TaskRepository
    from backend.repositories.user_repository import UserRepository
    from backend.services.auth_service import AuthService
    from backend.services.dashboard_service import DashboardService
    from backend.services.master_service import MasterService
    from backend.services.notification_service import NotificationService
    from backend.services.task_service import TaskService
    from backend.services.task_projection_service import TaskProjectionService
    from backend.services.user_service import UserService
    from backend.validators.uploads import (
        ATTACHMENT_RULES,
        IMAGE_SIGNATURES,
        IMAGE_SUFFIXES,
        MAX_IMAGE_BYTES,
        validate_attachment_upload,
        validate_image_upload,
    )

# Keep the original backend-root configuration boundary after moving this module
# beneath ``backend/app``. This preserves static and private-upload locations.
BASE_DIR = Path(__file__).resolve().parent.parent
SETTINGS = load_settings(BASE_DIR)
LOGGER = configure_logging(SETTINGS.log_level)
STATIC_DIR = SETTINGS.static_dir
# This directory is deliberately not mounted as static content. Files are served only
# after task-level authorization succeeds.
UPLOAD_DIR = SETTINGS.upload_dir
MAX_IMAGE_COUNT = 5
# Editing is intentionally bounded independently of compressed file size so a
# decompression-bomb image cannot exhaust the local application host.
MAX_IMAGE_PIXELS = 40_000_000
MAX_ATTACHMENT_COUNT = 10
SESSIONS: dict[str, tuple[int, datetime]] = {}
# Argon2id is deliberately the default (argon2-cffi's PasswordHasher type).  These
# parameters meet the interactive-password baseline while keeping the values in the
# encoded hash, so stronger future settings can trigger a transparent rehash.
PASSWORD_HASHER = PasswordHasher(time_cost=3, memory_cost=65536, parallelism=2, hash_len=32, salt_len=16)
SEOUL_TIMEZONE = ZoneInfo("Asia/Seoul")
TASK_REPOSITORY = TaskRepository()
DASHBOARD_REPOSITORY = DashboardRepository()
USER_REPOSITORY = UserRepository()
MASTER_REPOSITORY = MasterRepository()
AUDIT_REPOSITORY = AuditRepository()
BOOTSTRAP_REPOSITORY = BootstrapRepository()


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def today_in_seoul() -> date:
    """Return the calendar date used by the product's received/completed fields."""
    return datetime.now(SEOUL_TIMEZONE).date()


def legacy_hash_password(password: str, salt: bytes | None = None) -> str:
    """PBKDF2 format used only to verify and migrate pre-Argon2 local data."""
    salt = salt or secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), salt, 210_000)
    return f"{salt.hex()}${digest.hex()}"


def hash_password(password: str) -> str:
    """Create a non-reversible Argon2id password hash for every new password."""
    return PASSWORD_HASHER.hash(password)


def verify_password(password: str, encoded: str) -> bool:
    """Verify Argon2id hashes and the previous PBKDF2 format during migration."""
    try:
        if encoded.startswith("$argon2"):
            return PASSWORD_HASHER.verify(encoded, password)
        salt_hex, digest_hex = encoded.split("$", 1)
        actual = legacy_hash_password(password, bytes.fromhex(salt_hex)).split("$", 1)[1]
        return hmac.compare_digest(actual, digest_hex)
    except (InvalidHashError, VerificationError, ValueError, AttributeError, TypeError):
        return False


def password_hash_needs_upgrade(encoded: str) -> bool:
    """Identify PBKDF2 rows and Argon2 rows made with superseded parameters."""
    if not encoded.startswith("$argon2"):
        return True
    try:
        return PASSWORD_HASHER.check_needs_rehash(encoded)
    except (InvalidHashError, VerificationError):
        # Login verification will reject malformed hashes before this function runs.
        return False


@contextmanager
def db():
    """Compatibility boundary for routes and isolated tests; encryption lives in database.py."""
    connection = connect(DB_PATH)
    try:
        yield connection
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


def init_db() -> None:
    with db() as c:
        BOOTSTRAP_REPOSITORY.initialize(c, hash_password, now)


app, SCHEDULER = create_application(
    static_dir=STATIC_DIR, logger=LOGGER, initialize_database=init_db
)


def row_dict(row: sqlite3.Row | None):
    return dict(row) if row else None


def audit(c: sqlite3.Connection, actor: int | None, action: str, target_type: str, target_id: int | str, before=None, after=None, ip: str | None = None):
    import json
    AUDIT_REPOSITORY.create(
        c,
        actor,
        action,
        target_type,
        str(target_id),
        json.dumps(before, ensure_ascii=False) if before is not None else None,
        json.dumps(after, ensure_ascii=False) if after is not None else None,
        ip,
        now(),
    )


SENSITIVE_AUDIT_KEYS = {
    "password", "password_hash", "temporary_password", "current_password", "new_password", "token", "storage_key",
}


def audit_value(value):
    """Decode audit JSON for administrators without ever returning sensitive values."""
    import json
    if value is None:
        return None
    try:
        decoded = json.loads(value)
    except (TypeError, ValueError):
        return None

    def redact(item):
        if isinstance(item, dict):
            return {key: "[REDACTED]" if key.lower() in SENSITIVE_AUDIT_KEYS else redact(child) for key, child in item.items()}
        if isinstance(item, list):
            return [redact(child) for child in item]
        return item

    return redact(decoded)


def audit_log_view(row: sqlite3.Row) -> dict:
    """Expose useful change detail while omitting the raw stored JSON columns."""
    out = dict(row)
    out["before"] = audit_value(out.pop("before_data"))
    out["after"] = audit_value(out.pop("after_data"))
    return out


def current_user(request: Request) -> sqlite3.Row:
    token = request.headers.get("Authorization", "").removeprefix("Bearer ").strip()
    session = SESSIONS.get(token)
    if not session or session[1] <= datetime.now(timezone.utc):
        raise HTTPException(status_code=401, detail="로그인이 필요합니다.")
    with db() as c:
        user = USER_REPOSITORY.find_active_by_id(c, session[0])
    if not user:
        SESSIONS.pop(token, None)
        raise HTTPException(status_code=401, detail="사용할 수 없는 계정입니다.")
    return user


def admin_required(user: sqlite3.Row = Depends(current_user)) -> sqlite3.Row:
    user = password_changed_user(user)
    if user["role"] != "ADMIN":
        raise HTTPException(status_code=403, detail="관리자 권한이 필요합니다.")
    return user


def password_changed_user(user: sqlite3.Row) -> sqlite3.Row:
    """Prevent a temporary password from being used for work data access."""
    if user["must_change_password"]:
        raise HTTPException(status_code=403, detail="최초 로그인 후 비밀번호를 변경해야 합니다.")
    return user


app.include_router(
    create_auth_router(
        AuthService(
            db,
            verify_password,
            password_hash_needs_upgrade,
            hash_password,
            now,
            SESSIONS,
            audit,
            AuthRepository(),
        ),
        current_user,
    )
)
app.include_router(
    create_users_router(
        UserService(
            db,
            verify_password,
            hash_password,
            now,
            SESSIONS,
            audit,
            sqlite3.IntegrityError,
            USER_REPOSITORY,
        ),
        current_user,
        admin_required,
        password_changed_user,
    )
)


def accessible_task(c: sqlite3.Connection, task_id: int, user: sqlite3.Row) -> sqlite3.Row:
    task = TASK_REPOSITORY.live_task_by_id(c, task_id)
    if not task:
        raise HTTPException(status_code=404, detail="업무를 찾을 수 없습니다.")
    if user["role"] == "ADMIN" or task["created_by"] == user["id"]:
        return task
    if not TASK_REPOSITORY.task_has_assignee(c, task_id, user["id"]):
        raise HTTPException(status_code=403, detail="해당 업무에 접근할 권한이 없습니다.")
    return task



@app.get("/")
def index():
    return FileResponse(STATIC_DIR / "index.html")


def organization_view(row: sqlite3.Row) -> dict:
    out = dict(row)
    out["is_active"] = bool(out["is_active"])
    return out


def contact_view(c: sqlite3.Connection, row: sqlite3.Row) -> dict:
    out = dict(row)
    out["is_active"] = bool(out["is_active"])
    out["organization"] = MASTER_REPOSITORY.contact_organization_name(c, row["organization_id"])
    return out


master_service = MasterService(
    db, organization_view, contact_view, audit_log_view, now, audit, MASTER_REPOSITORY
)

def image_view(row: sqlite3.Row) -> dict:
    """Metadata deliberately excludes the private storage key."""
    return {
        "id": row["id"], "original_filename": row["original_filename"], "mime_type": row["mime_type"],
        "size_bytes": row["size_bytes"], "sort_order": row["sort_order"], "created_at": row["created_at"],
    }


def attachment_view(row: sqlite3.Row) -> dict:
    """Public attachment metadata; private paths and checksums never leave the API."""
    return {
        "id": row["id"], "original_filename": row["original_filename"], "mime_type": row["mime_type"],
        "attachment_type": row["attachment_type"], "size_bytes": row["size_bytes"], "created_at": row["created_at"],
    }


def private_upload_path(storage_key: str, missing_detail: str = "첨부 파일을 찾을 수 없습니다.") -> Path:
    root = UPLOAD_DIR.resolve()
    candidate = (root / storage_key).resolve()
    if root != candidate and root not in candidate.parents:
        raise HTTPException(status_code=404, detail=missing_detail)
    # Apply the Windows filesystem namespace only after resolving and enforcing
    # containment. Nested private uploads can exceed the legacy MAX_PATH limit.
    if os.name == "nt":
        native = str(candidate)
        if not native.startswith("\\\\?\\"):
            candidate = Path("\\\\?\\UNC\\" + native[2:] if native.startswith("\\\\") else "\\\\?\\" + native)
    return candidate


def private_image_path(storage_key: str) -> Path:
    return private_upload_path(storage_key, "이미지를 찾을 수 없습니다.")


def rotate_private_image(content: bytes, mime_type: str, direction: str) -> bytes:
    """Decode, rotate, and re-encode a private image within strict local limits.

    The source was MIME/signature checked at upload time, but decoding is still
    defensive because a compressed image can expand dramatically. Metadata is not
    copied into the derivative, avoiding accidental retention of EXIF location data.
    """
    formats = {"image/jpeg": "JPEG", "image/png": "PNG", "image/webp": "WEBP"}
    image_format = formats.get(mime_type)
    if not image_format:
        raise HTTPException(status_code=422, detail="편집할 수 없는 이미지 형식입니다.")
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(BytesIO(content)) as source:
                source.load()
                if source.width * source.height > MAX_IMAGE_PIXELS:
                    raise Image.DecompressionBombError("image dimensions exceed the editing limit")
                operation = Image.Transpose.ROTATE_270 if direction == "CLOCKWISE" else Image.Transpose.ROTATE_90
                edited = source.transpose(operation)
                if image_format == "JPEG" and edited.mode not in ("RGB", "L"):
                    edited = edited.convert("RGB")
                elif image_format == "WEBP" and edited.mode not in ("RGB", "RGBA"):
                    edited = edited.convert("RGBA" if "transparency" in source.info else "RGB")
                result = BytesIO()
                save_options = {"format": image_format}
                if image_format == "JPEG":
                    save_options.update({"quality": 90, "optimize": True})
                elif image_format == "WEBP":
                    save_options.update({"quality": 90, "method": 4})
                else:
                    save_options["optimize"] = True
                edited.save(result, **save_options)
                output = result.getvalue()
    except (Image.DecompressionBombError, Image.DecompressionBombWarning, UnidentifiedImageError, OSError, ValueError):
        raise HTTPException(status_code=422, detail="안전하게 편집할 수 없는 이미지입니다.")
    if not output or len(output) > MAX_IMAGE_BYTES or not IMAGE_SIGNATURES[mime_type](output):
        raise HTTPException(status_code=422, detail="편집 결과가 허용된 이미지 형식 또는 크기를 충족하지 않습니다.")
    return output


def task_view(c: sqlite3.Connection, task: sqlite3.Row) -> dict:
    out = dict(task)
    out["created_by_name"] = TASK_REPOSITORY.task_creator_name(c, task["created_by"])
    out["assignees"] = [dict(row) for row in TASK_REPOSITORY.task_assignee_rows(c, task["id"])]
    out["contact"] = row_dict(TASK_REPOSITORY.task_contact_row(c, task["requester_contact_id"]))
    out["images"] = [image_view(row) for row in TASK_REPOSITORY.task_image_rows(c, task["id"])]
    out["attachments"] = [attachment_view(row) for row in TASK_REPOSITORY.task_attachment_rows(c, task["id"])]
    return out


def task_changes(before: dict, after: dict) -> list[str]:
    labels = {
        "status": "상태", "description": "업무 설명", "received_date": "접수일",
        "due_date": "마감일", "completed_date": "처리일", "contact": "요청 정보", "assignees": "처리자",
    }
    return [labels[key] for key in labels if before.get(key) != after.get(key)]


def task_history(c: sqlite3.Connection, task_id: int) -> list[dict]:
    import json
    rows = TASK_REPOSITORY.task_history_rows(c, task_id)
    history = []
    for row in rows:
        item = dict(row)
        item["before"] = json.loads(item.pop("before_data")) if item["before_data"] else None
        item["after"] = json.loads(item.pop("after_data")) if item["after_data"] else None
        item["changed_fields"] = task_changes(item["before"], item["after"]) if item["before"] and item["after"] else []
        history.append(item)
    return history


def excel_text(value) -> str:
    """Return spreadsheet-safe text without allowing a task value to become a formula."""
    text = "" if value is None else str(value)
    return f"'{text}" if text.startswith(("=", "+", "-", "@")) else text


def status_label(task_status: str) -> str:
    return {"RECEIVED": "접수", "IN_PROGRESS": "진행중", "COMPLETED": "처리"}[task_status]


def task_export_workbook(c: sqlite3.Connection, rows, exported_by: sqlite3.Row, task_views: Callable[[Any, list[Any]], list[dict]]) -> bytes:
    """Build a local XLSX workbook containing only already-authorized task rows."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "업무 목록"
    headers = [
        "업무 ID", "상태", "업무 설명", "요청 담당자", "연락처", "소속 기관", "처리자",
        "접수일", "마감일", "처리일", "생성자", "생성 시각", "수정 시각",
    ]
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="2463A6")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    views = task_views(c, rows)
    for view in views:
        sheet.append([
            view["id"], status_label(view["status"]), excel_text(view["description"]),
            excel_text(view["contact"]["name"]), excel_text(view["contact"]["phone"]),
            excel_text(view["contact"]["organization"]),
            excel_text(", ".join(assignee["name"] for assignee in view["assignees"])),
            view["received_date"], view["due_date"] or "", view["completed_date"] or "",
            excel_text(view["created_by_name"]), view["created_at"], view["updated_at"],
        ])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [10, 12, 48, 18, 18, 24, 28, 14, 14, 14, 18, 28, 28]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    guide = workbook.create_sheet("내보내기 정보")
    guide.append(["항목", "값"])
    guide.append(["내보낸 시각(UTC)", now()])
    guide.append(["내보낸 사용자", excel_text(exported_by["name"])])
    guide.append(["업무 건수", len(rows)])
    guide.append(["보안", "이 파일은 로그인 사용자의 서버 권한 범위와 적용된 필터만 포함합니다."])
    for cell in guide[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["B"].width = 86
    for row in guide.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


# Register after the existing static task-export route so its literal path retains
# precedence over the legacy ``/api/tasks/{task_id}`` parameter route.
app.include_router(
    create_notification_router(
        NotificationService(
            db, TASK_REPOSITORY.authorized_task_query, now, lambda: today_in_seoul(), audit, NotificationRepository()
        ),
        current_user,
        password_changed_user,
    )
)

task_projection_service = TaskProjectionService(TASK_REPOSITORY, image_view, attachment_view)

app.include_router(
    create_dashboard_router(
        DashboardService(
            db, task_view, task_projection_service.task_views, task_export_workbook, audit, DASHBOARD_REPOSITORY
        ),
        current_user,
        password_changed_user,
    )
)
app.include_router(
    create_task_router(
        TaskService(
            db,
            task_view,
            task_projection_service.task_views,
            accessible_task,
            master_service.contact_from_input,
            lambda: today_in_seoul(),
            lambda: now(),
            audit,
            task_changes,
            task_history,
            private_image_path,
            private_upload_path,
            image_view,
            attachment_view,
            rotate_private_image,
            IMAGE_SUFFIXES,
            UPLOAD_DIR,
            MAX_IMAGE_COUNT,
            MAX_IMAGE_BYTES,
            validate_image_upload,
            ATTACHMENT_RULES,
            MAX_ATTACHMENT_COUNT,
            validate_attachment_upload,
            sqlite3.IntegrityError,
            TASK_REPOSITORY,
        ),
        current_user,
        password_changed_user,
        admin_required,
    )
)
app.include_router(
    create_master_router(
        master_service,
        current_user,
        password_changed_user,
        admin_required,
    )
)
