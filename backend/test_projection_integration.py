"""Public API and verification-runner regression for bulk task projections."""
from __future__ import annotations

from contextlib import contextmanager
from io import BytesIO
import importlib.util
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

import main


class ProjectionApiRegressionTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="calendar-projection-api-")
        self.addCleanup(temporary.cleanup)
        directory = Path(temporary.name)
        for name, value in (("DB_PATH", directory / "calendar.db"), ("UPLOAD_DIR", directory / "uploads")):
            override = patch.object(main, name, value)
            override.start()
            self.addCleanup(override.stop)
        override = patch.dict(os.environ, {"WORK_CALENDAR_INITIAL_ADMIN_PASSWORD": "projection-bootstrap-123"})
        override.start()
        self.addCleanup(override.stop)
        main.SESSIONS.clear()
        self.addCleanup(main.SESSIONS.clear)
        main.init_db()
        self.client = TestClient(main.app)
        self.admin = self.activate("admin", "projection-bootstrap-123")
        admin_id = self.client.get("/api/me", headers=self.admin).json()["id"]
        self.member_ids = []
        for login in ("projection_member_a", "projection_member_b"):
            response = self.client.post("/api/admin/users", headers=self.admin, json={
                "login_id": login, "name": login, "role": "MEMBER",
                "temporary_password": "projection-temporary-123",
            })
            self.assertEqual(response.status_code, 201, response.text)
            self.member_ids.append(response.json()["id"])
        self.member = self.activate("projection_member_a", "projection-temporary-123")
        self.tasks = []
        for index in range(20):
            status = ("RECEIVED", "IN_PROGRESS", "COMPLETED")[index % 3]
            response = self.client.post("/api/tasks", headers=self.admin, json={
                "assignee_ids": [self.member_ids[index % 2]],
                "contact": {"name": "projection-contact", "phone": "010-1234-5678", "organization": "projection-org"},
                "description": "=projection-fixture-" + str(index),
                "received_date": "2035-06-10", "due_date": "2035-06-15", "status": status,
                "completed_date": "2035-06-11" if status == "COMPLETED" else None,
            })
            self.assertEqual(response.status_code, 201, response.text)
            self.tasks.append(response.json())
        self.assertTrue(all(item["created_by"] == admin_id for item in self.tasks))

    def activate(self, login: str, password: str) -> dict[str, str]:
        response = self.client.post("/api/auth/login", json={"login_id": login, "password": password})
        self.assertEqual(response.status_code, 200, response.text)
        headers = {"Authorization": "Bearer " + response.json()["token"]}
        response = self.client.post("/api/me/password", headers=headers, json={
            "current_password": password, "new_password": "projection-active-password-123",
        })
        self.assertEqual(response.status_code, 200, response.text)
        return headers

    @contextmanager
    def count_selects(self):
        statements = []
        original = main.connect

        def traced_connect(path):
            connection = original(path)
            connection.set_trace_callback(
                lambda sql: statements.append(sql) if sql.lstrip().upper().startswith("SELECT ") else None
            )
            return connection

        with patch.object(main, "connect", side_effect=traced_connect):
            yield statements

    def test_list_query_budget_is_constant_and_matches_single_task_contract(self):
        counts = []
        for page_size in (1, 20):
            with self.count_selects() as selects:
                response = self.client.get(
                    f"/api/tasks?q=projection-fixture&page_size={page_size}", headers=self.admin,
                )
            self.assertEqual(response.status_code, 200, response.text)
            items = response.json()["items"]
            self.assertEqual(len(items), page_size)
            self.assertEqual(response.json()["pagination"]["total"], 20)
            counts.append(len(selects))
            for item in items:
                detail = self.client.get(f"/api/tasks/{item['id']}", headers=self.admin)
                self.assertEqual(detail.status_code, 200, detail.text)
                self.assertEqual(item, detail.json())
        self.assertEqual(counts[0], counts[1], counts)
        self.assertLessEqual(counts[1], 9, counts)
        print(f"API_LIST_SELECTS rows=1:{counts[0]} rows=20:{counts[1]}")

    def test_member_scope_is_preserved_in_list_and_both_calendars(self):
        allowed = {item["id"] for item in self.tasks if any(a["id"] == self.member_ids[0] for a in item["assignees"])}
        completed = {item["id"] for item in self.tasks if item["id"] in allowed and item["status"] == "COMPLETED"}
        with self.count_selects() as selects:
            response = self.client.get("/api/tasks?q=projection-fixture&page_size=100", headers=self.member)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual({item["id"] for item in response.json()["items"]}, allowed)
        self.assertLessEqual(len(selects), 9)
        for route, expected in (("/api/calendar?month=2035-06", allowed), ("/api/completed-calendar?month=2035-06", completed)):
            with self.count_selects() as selects:
                response = self.client.get(route, headers=self.member)
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual({item["id"] for item in response.json()}, expected)
            self.assertLessEqual(len(selects), 7)
        hidden = next(item["id"] for item in self.tasks if item["id"] not in allowed)
        self.assertEqual(self.client.get(f"/api/tasks/{hidden}", headers=self.member).status_code, 403)
        forbidden = self.client.get(f"/api/tasks?assignee_id={self.member_ids[1]}", headers=self.member)
        self.assertEqual(forbidden.status_code, 403)
        page_one = self.client.get("/api/tasks?q=projection-fixture&page_size=3&page=1", headers=self.member).json()
        page_two = self.client.get("/api/tasks?q=projection-fixture&page_size=3&page=2", headers=self.member).json()
        self.assertEqual(len(page_one["items"]), 3)
        self.assertEqual(len(page_two["items"]), 3)
        self.assertFalse({i["id"] for i in page_one["items"]} & {i["id"] for i in page_two["items"]})

    def test_scoped_xlsx_bulk_query_budget_and_formula_safe_fields(self):
        for headers in (self.admin, self.member):
            authorized = self.client.get("/api/tasks?q=projection-fixture&page_size=100", headers=headers).json()["items"]
            with self.count_selects() as selects:
                response = self.client.get("/api/tasks/export.xlsx?q=projection-fixture", headers=headers)
            self.assertEqual(response.status_code, 200, response.text)
            self.assertLessEqual(len(selects), 7)
            workbook = load_workbook(BytesIO(response.content))
            self.addCleanup(workbook.close)
            rows = list(workbook["업무 목록"].iter_rows(min_row=2, values_only=True))
            self.assertEqual({row[0] for row in rows}, {item["id"] for item in authorized})
            by_id = {item["id"]: item for item in authorized}
            for row in rows:
                item = by_id[row[0]]
                self.assertEqual(row[2], "'" + item["description"])
                self.assertEqual(row[3:6], (item["contact"]["name"], item["contact"]["phone"], item["contact"]["organization"]))
                self.assertEqual(row[6], ", ".join(a["name"] for a in item["assignees"]))
                self.assertEqual(row[8], item["due_date"])
                self.assertEqual(row[10], item["created_by_name"])
            print(f"API_XLSX_SELECTS rows={len(rows)} queries={len(selects)}")


class VerificationRunnerRegressionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        path = Path(__file__).resolve().parents[1] / "scripts" / "verify.py"
        spec = importlib.util.spec_from_file_location("calendar_verification", path)
        cls.runner = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.runner)

    def test_runner_overrides_product_settings_without_writing_a_plaintext_key(self):
        with tempfile.TemporaryDirectory() as temporary, patch.dict(os.environ, {
            "WORK_CALENDAR_DB": "DO_NOT_OPEN.db", "WORK_CALENDAR_DB_KEY_FILE": "DO_NOT_READ.dpapi",
            "WORK_CALENDAR_DB_KEY": "DO_NOT_USE", "WORK_CALENDAR_UNEXPECTED": "DO_NOT_KEEP",
        }):
            directory = Path(temporary)
            env = self.runner.isolated_environment(directory)
            self.assertEqual(Path(env["WORK_CALENDAR_DB"]).parent, directory)
            self.assertEqual(Path(env["WORK_CALENDAR_DB_KEY_FILE"]).parent, directory)
            self.assertNotIn("WORK_CALENDAR_UNEXPECTED", env)
            self.assertNotEqual(env["WORK_CALENDAR_DB_KEY"], "DO_NOT_USE")
            self.assertEqual(len(bytes.fromhex(env["WORK_CALENDAR_DB_KEY"])), 32)
            self.assertEqual(env["PYTHONNOUSERSITE"], "1")
            self.assertEqual(list(directory.iterdir()), [])

    def test_backend_gate_rejects_zero_or_skipped_tests(self):
        from subprocess import CompletedProcess
        for text in ("Ran 0 tests in 0.000s\nOK", "Ran 2 tests in 0.100s\nOK (skipped=1)"):
            with patch.object(self.runner, "run", return_value=CompletedProcess([], 0, "", text)):
                with self.assertRaises(RuntimeError):
                    self.runner.verify_backend({})

    def test_frontend_gate_checks_real_web_asset(self):
        from subprocess import CompletedProcess
        with patch.object(self.runner, "run", return_value=CompletedProcess([], 0, "", "")) as execute:
            self.runner.verify_frontend({})
        self.assertEqual(Path(execute.call_args.args[0][-1]), self.runner.ROOT / "web" / "app.js")


if __name__ == "__main__":
    unittest.main()
